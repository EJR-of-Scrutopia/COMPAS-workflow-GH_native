from __future__ import annotations

import sys
from pathlib import Path

import pytest

pytest.importorskip("fastapi")
sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "bench" / "studio"))

from fastapi.testclient import TestClient

import app as studio_app
import exports


@pytest.fixture()
def client(monkeypatch):
    # Setting the folder rebinds a module global, as for the other six; put
    # it back afterwards so no later test reads this one's temporary folder.
    monkeypatch.setattr(studio_app, "CABLENET_EXPORTS_DIR",
                        studio_app.CABLENET_EXPORTS_DIR, raising=False)
    return TestClient(studio_app.create_app(runner=lambda request: {}))


def test_the_export_folder_is_set_validated_and_remembered(client, tmp_path, monkeypatch):
    stored = {}
    monkeypatch.setattr(studio_app, "remember_setting",
                        lambda key, value: stored.__setitem__(key, value))
    monkeypatch.setattr(studio_app, "read_settings", lambda: dict(stored))

    chosen = tmp_path / "exports"
    chosen.mkdir()
    body = client.post("/api/cablenet/exports/folder",
                       json={"path": str(chosen)}).json()
    assert body["path"] == str(chosen)
    assert stored["cablenet_exports_folder"] == str(chosen)
    assert client.get("/api/cablenet/exports/folder").json()["path"] == str(chosen)


def test_a_folder_that_is_not_there_is_refused_by_name(client, tmp_path):
    missing = tmp_path / "nope"
    response = client.post("/api/cablenet/exports/folder",
                           json={"path": str(missing)})
    assert response.status_code == 400
    assert str(missing) in response.json()["detail"]


def test_an_empty_path_is_refused(client):
    assert client.post("/api/cablenet/exports/folder", json={"path": "  "}).status_code == 400


def test_browse_returns_a_path_without_setting_it(client, tmp_path, monkeypatch):
    seen = []
    monkeypatch.setattr(studio_app, "ask_for_folder", lambda *a, **k: str(tmp_path))
    monkeypatch.setattr(studio_app, "remember_setting",
                        lambda key, value: seen.append(key))
    assert client.post("/api/cablenet/exports/folder/browse").json()["path"] == str(tmp_path)
    assert seen == []        # browsing must not remember; the caller posts it back


def _demand(**kwargs):
    base = {
        "study": "Test Vault", "prestress": 300.0, "ea_newtons": 450000.0,
        "acceptance": 2.18, "acceptance_source": "plywood rib 2000 x 400",
        "sizing_stage": "S7", "density": 1800.0, "thickness": 0.02,
        "net": {"net_edge_count": 2253, "fixed": list(range(36)),
                "ea_provenance": "rope-4mm, assumed"},
        "wires": [{"name": "w1", "net_vertex": 658, "frame_point": [0, 0, 6000]}],
        "stages": [
            {"stage": 1, "name": "S1", "kind": "raise",
             "placed_weight_newtons": 0.0, "skin_load_sum_newtons": 0.0,
             "wire_rest_lengths": [2000.0], "wire_reel_commands": [0.0],
             "wire_tensions": [400.0], "deviation": 0.5, "reachable": True,
             "residual_after": 0.5},
            {"stage": 2, "name": "S7", "kind": "tile",
             "placed_weight_newtons": 12000.0, "skin_load_sum_newtons": 12000.0,
             "wire_rest_lengths": [1990.0], "wire_reel_commands": [-10.0],
             "wire_tensions": [900.0], "deviation": 1.4, "reachable": True,
             "residual_after": 1.4},
        ],
    }
    base.update(kwargs)
    return base


def _configuration(**kwargs):
    base = dict(motor="34HS46", drive="CL86Y", gearbox="EG23-G20", drum="drum-72",
                rope="rope-4mm", rail="MGN15H-300", sheave=None, reeve_factor=1,
                chain=["eye-M12", "turnbuckle-hook-hook-M10"])
    base.update(kwargs)
    return base


def _row(parts, configuration, **kwargs):
    import catalogue
    ceiling, binding = catalogue.ceiling_for(parts, configuration, 10.0)
    base = {
        "configuration": configuration, "ceiling": ceiling, "binding": binding,
        "passes": True, "passes_note": None, "margin": ceiling / 900.0,
        "price": catalogue.price_of(parts, configuration),
        "rope_speed_mm_s": catalogue.rope_speed(parts, configuration),
        "refused": None,
    }
    base.update(kwargs)
    return base


def test_the_model_carries_every_ceiling_term_with_the_part_that_sets_it():
    import catalogue
    parts = catalogue.load_parts()
    configuration = _configuration()
    model = exports.export_model(parts, _demand(), _row(parts, configuration),
                                 configuration, 10.0, "2026-10-07")

    names = [term["name"] for term in model["terms"]]
    assert "rope tension" in names and "motor torque" in names
    binding = [term for term in model["terms"] if term["binds"]]
    assert len(binding) == 1
    assert binding[0]["part_id"] == "turnbuckle-hook-hook-M10"
    assert round(binding[0]["newtons"]) == 1471
    assert model["verdict"]["tension"]["binding"] == binding[0]["part_id"]


def test_the_verdict_has_two_separate_halves():
    import catalogue
    parts = catalogue.load_parts()
    configuration = _configuration()
    model = exports.export_model(parts, _demand(), _row(parts, configuration),
                                 configuration, 10.0, "2026-10-07")
    assert set(model["verdict"]) == {"tension", "shape", "holds"}
    assert model["verdict"]["shape"]["worst_residual_mm"] == 1.4
    assert model["verdict"]["shape"]["acceptance_mm"] == 2.18
    assert model["verdict"]["shape"]["within"] is True
    assert model["verdict"]["holds"] is True


def test_parts_can_carry_it_while_the_net_misses_the_shape():
    # the case the panel got wrong before the first spec's final review
    import catalogue
    parts = catalogue.load_parts()
    configuration = _configuration()
    demand = _demand()
    demand["stages"][1]["residual_after"] = 7.4
    demand["stages"][1]["deviation"] = 7.4
    model = exports.export_model(parts, demand, _row(parts, configuration),
                                 configuration, 10.0, "2026-10-07")
    assert model["verdict"]["tension"]["passes"] is True
    assert model["verdict"]["shape"]["within"] is False
    assert model["verdict"]["shape"]["worst_stage"] == "S7"
    assert model["verdict"]["holds"] is False


def test_an_unreachable_stage_fails_the_shape_half():
    import catalogue
    parts = catalogue.load_parts()
    configuration = _configuration()
    demand = _demand()
    demand["stages"][1]["reachable"] = False
    model = exports.export_model(parts, demand, _row(parts, configuration),
                                 configuration, 10.0, "2026-10-07")
    assert model["verdict"]["shape"]["within"] is False
    assert model["verdict"]["holds"] is False


def test_every_chosen_part_arrives_with_its_provenance():
    import catalogue
    parts = catalogue.load_parts()
    configuration = _configuration()
    model = exports.export_model(parts, _demand(), _row(parts, configuration),
                                 configuration, 10.0, "2026-10-07")
    for part in model["parts"]:
        assert set(part) >= {"kind", "id", "model", "supplier", "part_number",
                             "unit_price", "vat", "price_seen", "confidence"}
    ids = {part["id"] for part in model["parts"]}
    assert "34HS46" in ids and "turnbuckle-hook-hook-M10" in ids


def test_a_refused_configuration_cannot_be_modelled():
    import catalogue
    parts = catalogue.load_parts()
    configuration = _configuration(motor="boatlift-1hp", drive="none")
    with pytest.raises(exports.ExportError, match="cannot be built"):
        exports.export_model(parts, _demand(),
                             {"refused": "family C", "configuration": configuration},
                             configuration, 10.0, "2026-10-07")


def _model(**overrides):
    """The one fixture every export test builds on."""
    import catalogue
    parts = catalogue.load_parts()
    configuration = overrides.pop("configuration", None) or _configuration()
    demand = overrides.pop("demand", None) or _demand()
    angle = overrides.pop("angle_degrees", 10.0)
    row = overrides.pop("row", None) or _row(parts, configuration)
    return exports.export_model(parts, demand, row, configuration, angle,
                                overrides.pop("generated_at", "2026-10-07"))


def test_exactly_one_term_binds_whatever_the_configuration():
    reeved = _configuration(sheave="WZ-11-K", reeve_factor=2)
    for configuration in (
        _configuration(),
        _configuration(chain=["eye-M12"]),
        _configuration(motor="23HS45", drive="CL57Y", gearbox="direct"),
        reeved,
    ):
        model = _model(configuration=configuration)
        assert sum(1 for term in model["terms"] if term["binds"]) == 1, configuration


def test_every_assumption_names_what_it_is():
    model = _model()
    whats = [item["what"] for item in model["assumptions"]]
    assert all(whats) and len(set(whats)) == len(whats)
    assert {"rope EA", "gearbox efficiency", "flat torque derate",
            "anchor angle"} <= set(whats)
    assert any("assumed" in w or "EA" in w for w in whats)
    assert model["not_checked"]


def test_the_workbook_has_the_six_sheets_in_order(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    model = _model()
    written = exports.write_spreadsheet(model, tmp_path, "Test Vault-cablenet")
    assert len(written) == 1 and written[0].suffix == ".xlsx"
    book = openpyxl.load_workbook(written[0])
    assert exports.SHEETS == ("Read this", "Chosen", "Parts", "Stages", "Ladder", "Hold")
    assert book.sheetnames == ["Read this", "Chosen", "Parts", "Stages", "Ladder", "Hold"]


def test_an_unpriced_line_marks_the_total_a_floor(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    model = _model()                       # the drum carries no price
    written = exports.write_spreadsheet(model, tmp_path, "x")
    book = openpyxl.load_workbook(written[0])
    text = " ".join(str(c.value) for row in book["Parts"].iter_rows()
                    for c in row if c.value is not None)
    assert "floor" in text.lower()
    assert "drum-72" in text


def test_without_openpyxl_it_writes_csvs_and_says_so(tmp_path, monkeypatch):
    monkeypatch.setattr(exports, "_openpyxl", None)
    written = exports.write_spreadsheet(_model(), tmp_path, "x")
    assert len(written) == 6
    assert all(path.suffix == ".csv" for path in written)
    assert exports.last_spreadsheet_note() and "openpyxl" in exports.last_spreadsheet_note()


def test_the_chosen_sheet_shows_both_halves_of_the_verdict(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    model = _model()
    model["verdict"]["shape"]["within"] = False
    written = exports.write_spreadsheet(model, tmp_path, "x")
    book = openpyxl.load_workbook(written[0])
    text = " ".join(str(c.value) for row in book["Chosen"].iter_rows()
                    for c in row if c.value is not None).lower()
    assert "tension" in text and "shape" in text


def test_the_csv_path_says_what_the_workbook_would(tmp_path, monkeypatch):
    # the rows are built once; the CSV is the same content, and it runs today
    import csv
    monkeypatch.setattr(exports, "_openpyxl", None)
    written = exports.write_spreadsheet(_model(), tmp_path, "x")
    assert [p.name for p in written] == [
        "x-" + name.lower().replace(" ", "-") + ".csv" for name in exports.SHEETS]
    def text(name):
        path = next(p for p in written if p.name.endswith(name + ".csv"))
        with open(path, newline="", encoding="utf-8") as handle:
            return " ".join(c for row in csv.reader(handle) for c in row)
    parts = text("parts")
    assert "floor" in parts.lower() and "drum-72" in parts
    chosen = text("chosen").lower()
    assert "tension" in chosen and "shape" in chosen


def test_the_note_is_reset_on_every_call(tmp_path, monkeypatch):
    with monkeypatch.context() as patch:
        patch.setattr(exports, "_openpyxl", None)
        exports.write_spreadsheet(_model(), tmp_path, "x")
        assert exports.last_spreadsheet_note()
    if exports._openpyxl is None:
        pytest.skip("needs openpyxl for the successful run")
    exports.write_spreadsheet(_model(), tmp_path, "x")
    assert exports.last_spreadsheet_note() is None


def _scored(parts, configuration, **kwargs):
    import catalogue
    ceiling, binding = catalogue.ceiling_for(parts, configuration, 10.0)
    row = {"configuration": configuration, "ceiling": ceiling, "binding": binding,
           "price": catalogue.price_of(parts, configuration), "refused": None}
    row.update(kwargs)
    return row


def _laddered(*variants):
    import catalogue
    parts = catalogue.load_parts()
    configuration = _configuration()
    rows = [_scored(parts, configuration)]
    for change in variants:
        if isinstance(change, str):
            rows.append({"configuration": dict(configuration, motor="boatlift-1hp"),
                         "refused": change})
        else:
            rows.append(_scored(parts, dict(configuration, **change)))
    return exports.export_model(parts, _demand(), _row(parts, configuration),
                                configuration, 10.0, "2026-10-07", ladder_rows=rows)


def test_a_rung_lists_only_the_keys_that_differ():
    model = _laddered({"chain": ["eye-M12"]})
    rung = [r for r in model["ladder"] if not r["is_chosen"]][0]
    assert list(rung["changes"]) == ["chain"]
    assert rung["ceiling_newtons"] > 0 and rung["binding"]


def test_the_chosen_rung_is_marked_with_no_changes():
    model = _laddered({"chain": ["eye-M12"]})
    chosen = [r for r in model["ladder"] if r["is_chosen"]]
    assert len(chosen) == 1 and chosen[0]["changes"] == {}


def test_a_refused_rung_survives_with_its_text():
    model = _laddered("family C")
    rung = model["ladder"][1]
    assert rung["refused"] == "family C"
    assert rung["ceiling_newtons"] is None and rung["binding"] is None


def test_no_ladder_rows_gives_an_empty_list():
    assert _model()["ladder"] == []


def test_the_ladder_sheet_renders_the_real_rungs():
    rows = exports._sheet_rows(_laddered({"chain": ["eye-M12"]}, "family C"))["Ladder"]
    assert len(rows) == 4 and rows[1][0] == "Chosen"
    assert "chain" in rows[2][1] and rows[3][4].startswith("Not buildable")


def test_the_diagram_marks_the_same_part_the_verdict_names(tmp_path):
    model = _model()
    svg = exports.write_diagram(model, tmp_path, "x").read_text(encoding="utf-8")
    assert svg.startswith("<?xml") or svg.lstrip().startswith("<svg")
    assert model["verdict"]["tension"]["binding"] in svg
    # the binding element is marked, and exactly one is
    assert svg.count('class="binds"') == 1


def test_the_diagram_prints_the_terms_it_was_given_and_computes_nothing():
    model = _model()
    for term in model["terms"]:
        term["newtons"] = 4242.0          # nonsense, but it is what was given
    svg = exports.diagram_svg(model)
    assert svg.count("4242") >= len(model["terms"])


def test_a_single_fall_draws_no_moving_block(tmp_path):
    model = _model()
    assert "moving block" not in exports.diagram_svg(model).lower()
    reeved = _model(configuration=_configuration(reeve_factor=2, sheave="WZ-11-K"))
    assert "moving block" in exports.diagram_svg(reeved).lower()


def test_the_diagram_escapes_what_it_is_given():
    model = _model()
    model["study"] = 'Vault & <script>"'
    svg = exports.diagram_svg(model)
    assert "<script>" not in svg and "&amp;" in svg


def test_escaping_does_not_double_escape():
    assert exports._esc('&<>"') == "&amp;&lt;&gt;&quot;"


def test_the_marked_element_moves_with_the_terms_not_with_a_recomputation():
    reeved = _model(configuration=_configuration(reeve_factor=2, sheave="WZ-11-K"))
    svg = exports.diagram_svg(reeved)
    assert svg.count('class="binds"') == 1
    marked = svg.split('class="binds"')[1].split("</g>")[0]
    assert "WZ-11-K" in marked


def test_the_diagram_names_the_net_band():
    svg = exports.diagram_svg(_model())
    assert "36 anchors" in svg and "658" in svg and "S7" in svg


def test_every_box_is_at_least_as_wide_as_its_longest_line():
    model = _model(configuration=_configuration(reeve_factor=2, sheave="WZ-11-K"))
    model["terms"][0]["part_id"] = "a-very-long-catalogue-identifier-for-a-part"
    by_name = {t["name"]: t for t in model["terms"]}
    for name, title in exports._PATH:
        if name not in by_name:
            continue
        lines = exports._box_lines(name, title, by_name[name], model["configuration"])
        longest = max(exports._estimated_width(t, s) for t, s, _ in lines)
        assert exports._box_width(lines) >= longest
    svg = exports.diagram_svg(model)
    import re
    width = int(re.search(r'<svg[^>]* width="(\d+)"', svg).group(1))
    assert 'viewBox="0 0 {} '.format(width) in svg


def test_the_net_band_counts_agree_with_their_nouns():
    one = exports.diagram_svg(_model())
    assert "1 wire," in one and "1 wires" not in one
    demand = _demand()
    demand["wires"].append({"name": "w2", "net_vertex": 9, "frame_point": [0, 0, 0]})
    demand["net"]["fixed"] = [1]
    two = exports.diagram_svg(_model(demand=demand))
    assert "2 wires" in two and "1 anchor," in two


def _section(text, heading_word):
    """The body of the `## ` section whose heading contains the word.

    Found by heading text, never by position, so adding a section cannot make
    an assertion silently read the wrong one.
    """
    blocks = text.split("\n## ")
    matches = [b for b in blocks[1:] if b.splitlines()[0].lower().find(heading_word) >= 0]
    assert len(matches) == 1, (heading_word, [b.splitlines()[0] for b in blocks])
    return matches[0]


def test_the_data_sheet_leads_with_the_failure_when_it_fails():
    model = _model()
    model["verdict"]["shape"]["within"] = False
    model["verdict"]["shape"]["worst_residual_mm"] = 7.4
    model["verdict"]["holds"] = False
    section = _section(exports.datasheet_markdown(model), "whether it holds")
    body = section.split("\n", 1)[1].strip()
    first_sentence = body.split(". ")[0]
    assert "not" in first_sentence.lower() or "fails" in first_sentence.lower()
    assert "does not keep its shape" in first_sentence
    assert "do not carry" not in first_sentence
    assert "7.4" in section


def test_the_data_sheet_names_the_tension_half_when_only_it_fails():
    model = _model()
    model["verdict"]["tension"]["passes"] = False
    model["verdict"]["holds"] = False
    section = _section(exports.datasheet_markdown(model), "whether it holds")
    first_sentence = section.split("\n", 1)[1].strip().split(". ")[0]
    assert "do not carry the tension" in first_sentence
    assert "does not keep" not in first_sentence


def test_the_data_sheet_states_a_pass_first_when_both_halves_hold():
    model = _model()
    assert model["verdict"]["holds"]
    section = _section(exports.datasheet_markdown(model), "whether it holds")
    first_sentence = section.split("\n", 1)[1].strip().split(". ")[0]
    assert "holds" in first_sentence.lower() and "not" not in first_sentence.lower()


def test_the_assumptions_are_a_section_and_not_a_footnote():
    text = exports.datasheet_markdown(_model())
    headings = [line for line in text.splitlines() if line.startswith("## ")]
    assert any("assumption" in h.lower() for h in headings)
    assert any("not checked" in h.lower() for h in headings)
    assert "450000" in text or "450,000" in text      # the rope EA is named


def test_the_sheet_has_the_nine_sections_in_the_specified_order():
    text = exports.datasheet_markdown(_model())
    headings = [l[3:].lower() for l in text.splitlines() if l.startswith("## ")]
    words = ["replaces", "demands", "chosen", "holds", "hold the weight",
             "grab the net", "load path", "assumptions", "not checked"]
    assert len(headings) == 9
    for heading, word in zip(headings, words):
        assert word in heading


def test_every_assumed_figure_reaches_the_sheet():
    model = _model()
    section = _section(exports.datasheet_markdown(model), "assumptions")
    for assumption in model["assumptions"]:
        assert assumption["what"] in section


def test_every_unchecked_item_reaches_the_sheet():
    model = _model()
    section = _section(exports.datasheet_markdown(model), "not checked")
    for line in model["not_checked"]:
        assert line in section


def test_the_acceptance_source_is_quoted_verbatim():
    model = _model()
    section = _section(exports.datasheet_markdown(model), "replaces")
    assert model["verdict"]["shape"]["acceptance_source"] in section


def test_the_data_sheet_has_no_em_dash_and_no_latex():
    text = exports.datasheet_markdown(_model())
    assert "\u2014" not in text and "$" not in text and "\frac" not in text


def test_the_load_path_lists_every_term_and_marks_the_binding_one():
    model = _model()
    section = _section(exports.datasheet_markdown(model), "load path")
    for term in model["terms"]:
        assert term["name"] in section
    assert section.count("binds") + section.count("binding") >= 1


def test_the_data_sheet_is_written_beside_the_others(tmp_path):
    path = exports.write_datasheet(_model(), tmp_path, "vault-cablenet")
    assert path.name == "vault-cablenet.md"
    assert path.read_text(encoding="utf-8").startswith("# ")


def test_the_model_names_the_units_of_density_and_thickness():
    demand = _model()["demand"]
    assert demand["density_kg_m3"] == 1800.0
    assert demand["thickness_m"] == 0.02


def test_the_sheet_names_the_units_of_density_and_thickness():
    section = _section(exports.datasheet_markdown(_model()), "demands")
    assert "1,800 kg/m3" in section and "0.020 m" in section
    assert "own units" not in section


def test_no_assumption_has_a_confidence_word_for_a_value():
    model = _model()
    for assumption in model["assumptions"]:
        assert assumption["value"] != "assumed", assumption
    text = exports.datasheet_markdown(model)
    assert "Value: assumed" not in text


def test_the_rope_ea_and_the_gearbox_efficiency_each_appear_once():
    whats = [a["what"] for a in _model()["assumptions"]]
    assert whats.count("rope EA") == 1
    assert whats.count("gearbox efficiency") == 1
    assert not any(w.endswith("ea confidence") or w.endswith("efficiency confidence")
                   for w in whats)


def test_the_load_path_uses_one_decimal_convention():
    import re
    section = _section(exports.datasheet_markdown(_model()), "load path")
    figures = re.findall(r"allows ([\d,.]+) N", section)
    assert figures and all(re.fullmatch(r"\d+\.\d", f) for f in figures)


# ---------------------------------------------------------------------------
# The routes
# ---------------------------------------------------------------------------

def _plant_demand(monkeypatch, tmp_path, name="My Vault", demand=None):
    import json
    import bundle
    import geometry
    from conftest_data import tiny_contract

    # The demand is found by the cut the study would be run with, which is read
    # from the study's own contract, so the study has to be there.
    upload = tmp_path / "upload"
    upload.mkdir(exist_ok=True)
    (upload / "{}-contract.json".format(name)).write_text(
        json.dumps(tiny_contract()), encoding="utf-8")
    monkeypatch.setattr(bundle, "UPLOAD_DIR", upload)
    monkeypatch.setattr(bundle, "STUDIES_DIR", tmp_path / "studies")
    path = bundle.cablenet_path(geometry.slugify(name), "tile", "herringbone",
                                1.0, 0.02, None)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(demand or _demand()), encoding="utf-8")


def _into(monkeypatch, folder):
    folder.mkdir(exist_ok=True)
    monkeypatch.setattr(studio_app, "read_settings",
                        lambda: {"cablenet_exports_folder": str(folder)})


def _spy_on_export_model(monkeypatch):
    seen = {}
    real = exports.export_model

    def spy(*args, **kwargs):
        seen["row"] = args[2]
        seen["ladder"] = kwargs.get("ladder_rows")
        return real(*args, **kwargs)

    monkeypatch.setattr(exports, "export_model", spy)
    return seen


def test_one_run_produces_three_documents_that_name_the_same_binding_part(tmp_path):
    # LITERAL: the ceiling is rendered once by the shared formatter and must
    # appear byte for byte in the diagram, the data sheet and the spreadsheet.
    model = _model()
    paths = (exports.write_spreadsheet(model, tmp_path, "x")
             + [exports.write_diagram(model, tmp_path, "x"),
                exports.write_datasheet(model, tmp_path, "x")])
    binding = model["verdict"]["tension"]["binding"]
    shown = exports._newtons(model["verdict"]["tension"]["ceiling_newtons"])
    assert shown == "1471.0"
    svg = [p for p in paths if p.suffix == ".svg"][0].read_text(encoding="utf-8")
    md = [p for p in paths if p.suffix == ".md"][0].read_text(encoding="utf-8")
    assert binding in svg and binding in md
    assert shown + " N" in svg and shown + " N" in md
    if exports._openpyxl is not None:
        import openpyxl
        book = openpyxl.load_workbook(next(p for p in paths if p.suffix == ".xlsx"))
        cells = [c for row in book["Chosen"].iter_rows() for c in row
                 if c.value is not None]
        ceiling = next(c for c in cells if c.value == float(shown))
        assert ceiling.number_format == "0.0"
        assert binding in [c.value for c in cells]


def test_the_csv_fallback_prints_the_same_ceiling_text(tmp_path, monkeypatch):
    monkeypatch.setattr(exports, "_openpyxl", None)
    model = _model()
    paths = exports.write_spreadsheet(model, tmp_path, "x")
    shown = exports._newtons(model["verdict"]["tension"]["ceiling_newtons"])
    chosen = next(p for p in paths if p.name.endswith("-chosen.csv"))
    text = chosen.read_text(encoding="utf-8-sig")
    assert shown in text
    assert model["verdict"]["tension"]["binding"] in text


def test_every_force_is_written_by_the_one_formatter():
    # no renderer may format newtons its own way: 1500 N and 1471.4 N read alike
    assert exports._newtons(1500) == "1500.0"
    assert exports._newtons(1471.4) == "1471.4"
    assert exports._newtons(None) == "not recorded"
    model = _model()
    svg = exports.diagram_svg(model)
    md = exports.datasheet_markdown(model)
    import re
    for text in (svg, md):
        assert not re.search(r"\b\d{4,} N\b", text)       # always a decimal
    # no force is grouped, whatever follows it: the only grouped figure the
    # sheet may carry is the density, which is not a force
    assert _grouped_figures(md + svg) == [], _grouped_figures(md + svg)


GROUPED = r"\d{1,3}(?:,\d{3})+(?:\.\d+)?(?: \S+)?"


def _grouped_figures(text):
    import re
    return [g for g in re.findall(GROUPED, text) if not g.endswith(" kg/m3")]


def test_the_forbidding_tests_really_fail_on_a_grouped_force():
    # the guard must bite: feed it the shapes the old regexes let through
    import re
    for bad in ("Value: 450,000", "Value: 450,000.0", "allows 1,471.0 N",
                "taken as 450,000 N"):
        assert _grouped_figures(bad), bad
    assert _grouped_figures("a density of 1,800 kg/m3") == []
    assert not re.fullmatch(r"\d+\.\d", "1,471.0")


def test_a_force_assumption_prints_ungrouped_with_one_decimal_and_a_unit():
    text = exports.datasheet_markdown(_model())
    assert "**rope EA.** Value: 450000.0 N." in text
    assert "**uniform prestress.** Value: 300.0 N." in text
    assert "450,000" not in text


def test_the_route_writes_all_three_and_they_agree(client, monkeypatch, tmp_path):
    _plant_demand(monkeypatch, tmp_path)
    folder = tmp_path / "out"
    _into(monkeypatch, folder)
    response = client.post("/api/studies/My Vault/cablenet/exports",
                           json={"configuration": _configuration()})
    assert response.status_code == 200, response.text
    written = [Path(p) for p in response.json()["paths"]]
    assert {p.suffix for p in written} >= {".svg", ".md"}
    assert all(p.is_file() and p.parent == folder for p in written)
    svg = next(p for p in written if p.suffix == ".svg").read_text(encoding="utf-8")
    md = next(p for p in written if p.suffix == ".md").read_text(encoding="utf-8")
    assert "turnbuckle-hook-hook-M10" in svg and "turnbuckle-hook-hook-M10" in md
    assert "1471.0 N" in svg and "1471.0 N" in md


def test_the_route_scores_the_panels_ladder(client, monkeypatch, tmp_path):
    _plant_demand(monkeypatch, tmp_path)
    _into(monkeypatch, tmp_path / "out")
    seen = _spy_on_export_model(monkeypatch)
    response = client.post("/api/studies/My Vault/cablenet/exports",
                           json={"configuration": _configuration()})
    assert response.status_code == 200, response.text
    ladder = seen["ladder"]
    assert len(ladder) == 5
    assert ladder[0]["configuration"] == _configuration()
    assert ladder[1]["configuration"]["chain"] == ["eye-M12", "turnbuckle-eye-eye-M10"]
    assert ladder[2]["configuration"]["rope"] == "rope-5mm"
    assert ladder[3]["configuration"]["chain"] == ["eye-M16", "turnbuckle-eye-eye-M10"]
    assert ladder[3]["configuration"]["rope"] == "rope-6mm"
    assert ladder[4]["configuration"]["chain"] == ["eye-M20", "turnbuckle-eye-eye-M12"]
    assert ladder[4]["configuration"]["rope"] == "rope-8mm"
    assert all(rung["refused"] is None and rung["ceiling"] > 0 for rung in ladder)


def test_the_scored_rows_carry_the_drive_and_the_load_factor(client, monkeypatch, tmp_path):
    # the documents read the row, so the exports score with the demand in hand
    sizing = {"stage": "S7", "worst_wire_tension_newtons": 900.0,
              "worst_actuator_newtons": 0.0, "worst_sag_mm": 1.4, "load_newtons": 12000.0}
    _plant_demand(monkeypatch, tmp_path, demand=_demand(sizing=sizing))
    _into(monkeypatch, tmp_path / "out")
    seen = _spy_on_export_model(monkeypatch)
    response = client.post("/api/studies/My Vault/cablenet/exports",
                           json={"configuration": _configuration()})
    assert response.status_code == 200, response.text
    for row in [seen["row"]] + seen["ladder"]:
        assert row["drive"] == "CL86Y"
        assert row["load_factor"]["stage"] == "S7"
    # 1471 N against 900 N per unit skin: 1.6 passes, 1.7 breaches
    assert seen["row"]["load_factor"]["limit_factor"] == pytest.approx(1.6)
    assert seen["row"]["load_factor"]["binding"] == "anchor"


def test_the_scoring_route_and_the_exports_score_a_configuration_alike(
        client, monkeypatch, tmp_path):
    import json
    sizing = {"stage": "S7", "worst_wire_tension_newtons": 900.0,
              "worst_actuator_newtons": 0.0, "worst_sag_mm": 1.4, "load_newtons": 12000.0}
    _plant_demand(monkeypatch, tmp_path, demand=_demand(sizing=sizing))
    _into(monkeypatch, tmp_path / "out")
    seen = _spy_on_export_model(monkeypatch)
    scored = client.post("/api/studies/My Vault/cablenet/configurations", json={
        "configurations": [_configuration()], "angle_degrees": 10.0,
        "options": {"material": "tile", "pattern": "herringbone", "size": 1.0,
                    "thickness": 0.02}}).json()["rows"][0]
    exported = client.post("/api/studies/My Vault/cablenet/exports",
                           json={"configuration": _configuration(), "angle_degrees": 10.0})
    assert exported.status_code == 200, exported.text
    # one row, built once: the scoring route adds the note on the load factor and
    # nothing else to what the exports are given
    assert scored.pop("load_factor_note") is None
    assert scored == json.loads(json.dumps(seen["row"]))


def test_a_rung_the_catalogue_refuses_is_passed_through_with_its_text(
        client, monkeypatch, tmp_path):
    _plant_demand(monkeypatch, tmp_path)
    _into(monkeypatch, tmp_path / "out")
    seen = _spy_on_export_model(monkeypatch)
    # no eye-and-eye turnbuckle exists at M20, so the first rung is refused
    configuration = _configuration(chain=["eye-M20", "turnbuckle-hook-eye-M20"])
    response = client.post("/api/studies/My Vault/cablenet/exports",
                           json={"configuration": configuration})
    assert response.status_code == 200, response.text
    assert len(seen["ladder"]) == 5
    assert seen["ladder"][1]["refused"]


def test_the_rope_wound_is_the_worst_wires_summed_reel_commands(
        client, monkeypatch, tmp_path):
    stages = _demand()["stages"]
    stages[0]["wire_reel_commands"] = [5.0, -3.0]
    stages[1]["wire_reel_commands"] = [-10.0, -20.0]
    _plant_demand(monkeypatch, tmp_path, demand=_demand(stages=stages))
    _into(monkeypatch, tmp_path / "out")
    seen = _spy_on_export_model(monkeypatch)
    client.post("/api/studies/My Vault/cablenet/exports",
                json={"configuration": _configuration()})
    assert seen["row"]["rope_path"]["rope_wound_mm"] == 23.0


def test_a_body_rope_wound_wins_over_the_demand(client, monkeypatch, tmp_path):
    _plant_demand(monkeypatch, tmp_path)
    _into(monkeypatch, tmp_path / "out")
    seen = _spy_on_export_model(monkeypatch)
    client.post("/api/studies/My Vault/cablenet/exports",
                json={"configuration": _configuration(), "rope_wound_mm": 123.0})
    assert seen["row"]["rope_path"]["rope_wound_mm"] == 123.0


def test_a_failing_configuration_is_still_exported(client, monkeypatch, tmp_path):
    stages = _demand()["stages"]
    stages[1]["wire_tensions"] = [90000.0]       # far over any ceiling
    _plant_demand(monkeypatch, tmp_path, demand=_demand(stages=stages))
    _into(monkeypatch, tmp_path / "out")
    seen = _spy_on_export_model(monkeypatch)
    response = client.post("/api/studies/My Vault/cablenet/exports",
                           json={"configuration": _configuration()})
    assert response.status_code == 200, response.text
    assert seen["row"]["passes"] is False


def test_a_refused_configuration_is_a_400_not_a_500(client, monkeypatch, tmp_path):
    _plant_demand(monkeypatch, tmp_path)
    _into(monkeypatch, tmp_path / "out")
    response = client.post(
        "/api/studies/My Vault/cablenet/exports",
        json={"configuration": _configuration(motor="no-such-motor")})
    assert response.status_code == 400
    assert "cannot be built" in response.json()["detail"]


def test_a_study_with_no_demand_document_cannot_export(client):
    response = client.post("/api/studies/nothing-here/cablenet/exports",
                           json={"configuration": _configuration()})
    assert response.status_code == 404
    assert ("Upload the export, then press Run cable net analysis in the Cable net "
            "section and the engine will write one.") in response.json()["detail"]


def test_an_unwritable_folder_is_reported_and_not_swallowed(client, monkeypatch, tmp_path):
    _plant_demand(monkeypatch, tmp_path)
    gone = tmp_path / "gone"
    monkeypatch.setattr(studio_app, "read_settings",
                        lambda: {"cablenet_exports_folder": str(gone)})
    # the folder is missing, so the export must say so rather than silently
    # landing somewhere else, which is what deliver_output would do
    response = client.post("/api/studies/My Vault/cablenet/exports",
                           json={"configuration": _configuration()})
    assert response.status_code != 200
    assert str(gone) in response.json()["detail"]
    assert not gone.exists()


def test_an_os_error_while_writing_names_the_folder_and_the_reason(
        client, monkeypatch, tmp_path):
    _plant_demand(monkeypatch, tmp_path)
    folder = tmp_path / "out"
    _into(monkeypatch, folder)

    def refuse(*args, **kwargs):
        raise PermissionError(13, "Access is denied")

    monkeypatch.setattr(exports, "write_diagram", refuse)
    response = client.post("/api/studies/My Vault/cablenet/exports",
                           json={"configuration": _configuration()})
    assert response.status_code != 200
    detail = response.json()["detail"]
    assert str(folder) in detail and "Access is denied" in detail


def test_each_kind_can_be_downloaded_after_a_run(client, monkeypatch, tmp_path):
    _plant_demand(monkeypatch, tmp_path)
    _into(monkeypatch, tmp_path / "out")
    client.post("/api/studies/My Vault/cablenet/exports",
                json={"configuration": _configuration()})
    for kind, suffix in (("spreadsheet", ".xlsx"), ("diagram", ".svg"),
                         ("datasheet", ".md")):
        response = client.get("/api/studies/My Vault/cablenet/exports/" + kind)
        assert response.status_code == 200, kind
        assert suffix in response.headers["content-disposition"]
    assert client.get("/api/studies/My Vault/cablenet/exports/other").status_code == 400


def test_a_csv_fallback_is_served_as_one_zip(client, monkeypatch, tmp_path):
    import io
    import zipfile
    monkeypatch.setattr(exports, "_openpyxl", None)
    _plant_demand(monkeypatch, tmp_path)
    _into(monkeypatch, tmp_path / "out")
    posted = client.post("/api/studies/My Vault/cablenet/exports",
                         json={"configuration": _configuration()}).json()
    assert posted["note"]
    response = client.get("/api/studies/My Vault/cablenet/exports/spreadsheet")
    assert response.status_code == 200
    names = zipfile.ZipFile(io.BytesIO(response.content)).namelist()
    assert len(names) == len(exports.SHEETS) and all(n.endswith(".csv") for n in names)


def test_downloading_before_any_run_is_a_404(client, monkeypatch, tmp_path):
    _into(monkeypatch, tmp_path / "out")
    response = client.get("/api/studies/My Vault/cablenet/exports/diagram")
    assert response.status_code == 404


# ---------------------------------------------------------------------------
# The shape verdict has three states. Nothing to judge is NOT a pass.
# ---------------------------------------------------------------------------

def _degenerate_demands():
    no_stages = _demand(stages=[])
    no_residual = _demand()
    for stage in no_residual["stages"]:
        stage.pop("residual_after")
    no_line = _demand(acceptance=None)
    no_line["stages"][1]["residual_after"] = 99.0
    return {"no stages": no_stages, "no residual": no_residual,
            "no acceptance": no_line}


@pytest.mark.parametrize("which", ["no stages", "no residual", "no acceptance"])
def test_a_demand_with_nothing_to_judge_is_not_established(which):
    model = _model(demand=_degenerate_demands()[which])
    shape = model["verdict"]["shape"]
    assert shape["within"] is None
    assert shape["why_unknown"]
    assert model["verdict"]["holds"] is not True
    assert model["verdict"]["holds"] is None


def test_the_three_reasons_are_told_apart():
    reasons = {k: _model(demand=d)["verdict"]["shape"]["why_unknown"]
               for k, d in _degenerate_demands().items()}
    assert len(set(reasons.values())) == 3


def test_a_failing_tension_half_is_false_even_when_the_shape_is_unknown():
    import catalogue
    parts = catalogue.load_parts()
    configuration = _configuration()
    row = _row(parts, configuration, passes=False)
    model = _model(demand=_demand(stages=[]), row=row)
    assert model["verdict"]["holds"] is False


def test_genuine_verdicts_are_not_turned_into_unknowns():
    passing = _model()["verdict"]
    assert passing["shape"]["within"] is True and passing["holds"] is True
    assert passing["shape"]["why_unknown"] is None
    demand = _demand()
    demand["stages"][1]["residual_after"] = 7.4
    failing = _model(demand=demand)["verdict"]
    assert failing["shape"]["within"] is False and failing["holds"] is False


def test_an_unreachable_stage_is_a_failure_and_not_an_unknown():
    demand = _demand(stages=[{"stage": 1, "name": "S1", "reachable": False}])
    verdict = _model(demand=demand)["verdict"]
    assert verdict["shape"]["within"] is False
    assert verdict["holds"] is False
    nulled = _demand(acceptance=None)
    nulled["stages"][1]["reachable"] = False
    assert _model(demand=nulled)["verdict"]["shape"]["within"] is False


def _unknown_model():
    return _model(demand=_demand(stages=[]))


def test_the_data_sheet_says_not_established_and_never_that_it_holds():
    text = exports.datasheet_markdown(_unknown_model())
    section = _section(text, "whether it holds")
    first = section.split("\n", 1)[1].strip().split(". ")[0].lower()
    assert "not been established" in first
    assert "this configuration holds" not in section.lower()
    assert "does not hold" not in section.lower()
    assert "stays within the acceptance line. That has not been established" in section
    assert "no residual was measured" in section


def test_the_chosen_sheet_says_not_established_and_never_yes_for_the_shape(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    written = exports.write_spreadsheet(_unknown_model(), tmp_path, "x")
    rows = [[c.value for c in row] for row in
            openpyxl.load_workbook(written[0])["Chosen"].iter_rows()]
    by_label = {r[0]: r for r in rows if r and r[0]}
    shape_row = by_label["Shape: within the acceptance line"]
    both_row = by_label["Both halves hold"]
    assert shape_row[1] == "not established"
    assert "no residual was measured" in str(shape_row[2])
    assert both_row[1] == "not established"


def test_the_diagram_says_not_established_and_never_that_it_holds():
    # every case is sized, so the load factor line has a figure to give and
    # whatever the diagram says is unknown, it says of the verdict alone
    svg = exports.diagram_svg(_sized_model(stages=[]))
    assert "Verdict: not established" in svg
    assert "Verdict: holds" not in svg and "does not hold" not in svg
    passing = exports.diagram_svg(_sized_model())
    assert "Verdict: holds" in passing and "not established" not in passing
    demand = _v2_demand()
    demand["stages"][1]["residual_after"] = 7.4
    failing = exports.diagram_svg(_sized_model(stages=demand["stages"]))
    assert "does not hold" in failing and "not established" not in failing


def test_the_panel_and_the_documents_agree_on_what_is_unknown():
    import re
    js = (Path(__file__).resolve().parents[2] / "bench" / "studio" / "static"
          / "cablenet_model.js").read_text(encoding="utf-8")
    for demand in _degenerate_demands().values():
        why = _model(demand=demand)["verdict"]["shape"]["why_unknown"]
        assert why in js          # the panel carries the documents' wording
    assert re.search(r"acceptance === null", js)


# ---------------------------------------------------------------------------
# Review fixes: one truth test for `reachable`, one rounding for lengths
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("missing", ["absent", "null"])
def test_a_stage_with_no_reachable_is_told_the_same_way_everywhere(
        tmp_path, monkeypatch, missing):
    demand = _demand()
    if missing == "absent":
        demand["stages"][1].pop("reachable")
    else:
        demand["stages"][1]["reachable"] = None
    model = _model(demand=demand)
    # not False, so not an unreachable stage; and never read as reached
    assert model["verdict"]["shape"]["unreachable_stages"] == []
    text = exports.datasheet_markdown(model)
    assert "Stage 2 (S7, tile): residual 1.40 mm, reachability not recorded." in text
    assert "Stage 1 (S1, raise): residual 0.50 mm, reached." in text
    assert "not reachable" not in text
    rows = exports._sheet_rows(model)["Stages"]
    reach = {r[1]: r[5] for r in rows if len(r) == 7 and r[1] in ("S1", "S7")}
    assert reach == {"S1": "yes", "S7": "not recorded"}
    assert "Shape: unreachable stages" not in str(exports._sheet_rows(model)["Chosen"])


def test_a_stage_that_is_false_is_unreachable_in_every_document():
    demand = _demand()
    demand["stages"][1]["reachable"] = False
    model = _model(demand=demand)
    assert model["verdict"]["shape"]["unreachable_stages"] == ["S7"]
    assert "residual 1.40 mm, not reachable." in exports.datasheet_markdown(model)
    rows = exports._sheet_rows(model)["Stages"]
    assert {r[1]: r[5] for r in rows if len(r) == 7}["S7"] == "no"
    assert "does not hold" in exports.diagram_svg(model)


def test_the_panel_tests_reachable_against_false_only():
    js = (Path(__file__).resolve().parents[2] / "bench" / "studio" / "static"
          / "cablenet_model.js").read_text(encoding="utf-8")
    assert "stage.reachable === false" in js
    assert "stage.reachable ===" not in js.replace("stage.reachable === false", "")


def test_workbook_lengths_use_the_data_sheets_rounding(tmp_path, monkeypatch):
    demand = _demand()
    demand["stages"][1]["residual_after"] = 1.4000000000000001
    demand["stages"][1]["deviation"] = 0.30000000000000004
    monkeypatch.setattr(exports, "_openpyxl", None)
    model = _model(demand=demand)
    paths = exports.write_spreadsheet(model, tmp_path, "x")
    stages = next(p for p in paths if p.name.endswith("-stages.csv")).read_text(
        encoding="utf-8-sig")
    assert "1.40" in stages and "0.30" in stages
    assert "1.4000000000000001" not in stages and "0.30000000000000004" not in stages
    assert exports._millimetres(1.4000000000000001) == "1.40"


# ---------------------------------------------------------------------------
# A chosen rope can contradict the analysis the demand was solved for
# ---------------------------------------------------------------------------

def _all_documents(model, tmp_path):
    """The three documents as text, the workbook through its CSV fallback."""

    import re
    texts = {"datasheet": exports.datasheet_markdown(model),
             "diagram": exports.diagram_svg(model)}
    rows = exports._sheet_rows(model)
    texts["workbook"] = "\n".join(
        str(c) for sheet in rows.values() for row in sheet for c in row)
    return texts


def test_a_different_rope_is_stated_in_all_three_documents(tmp_path):
    # solved at 450000 N for rope-4mm, exported with rope-8mm (EA 1800000 N)
    model = _model(configuration=_configuration(rope="rope-8mm"))
    mismatch = model["rope_mismatch"]
    assert mismatch["analysed_rope"] == "rope-4mm"
    assert mismatch["chosen_rope"] == "rope-8mm"
    for name, text in _all_documents(model, tmp_path).items():
        flat = " ".join(text.split())
        assert "the chosen rope is not the rope that was analysed" in flat, name
        assert "rope-4mm" in flat and "rope-8mm" in flat, name
        assert "The tension ceiling is for the chosen rope" in flat, name
        assert ("The prestress floor, the residuals and the cut lengths are "
                "for the analysed rope and do not describe the chosen one"
                in flat), name
    # the ceiling really is the chosen rope's, and the export is not refused
    assert model["verdict"]["tension"]["ceiling_newtons"] > 0


def test_the_data_sheet_states_the_mismatch_before_the_first_section():
    text = exports.datasheet_markdown(
        _model(configuration=_configuration(rope="rope-8mm")))
    assert text.index("not the rope that was analysed") < text.index("## ")


def test_the_two_ea_figures_are_each_labelled_when_the_rope_differs():
    text = exports.datasheet_markdown(
        _model(configuration=_configuration(rope="rope-8mm")))
    assert "(EA 450000.0 N)" in text and "(EA 1800000.0 N)" in text


def test_nothing_extra_is_said_when_the_rope_is_the_one_analysed(tmp_path):
    import re
    model = _model()
    assert model["rope_mismatch"] is None
    for name, text in _all_documents(model, tmp_path).items():
        assert "not the rope that was analysed" not in text, name
        assert "Warning" not in text, name
    md = exports.datasheet_markdown(model)
    # ONE rope EA figure, wherever it is printed
    figures = set(re.findall(r"\b(?:450000|1800000)\.0\b", md))
    assert figures == {"450000.0"}


def test_a_demand_with_no_ea_cannot_contradict_a_rope():
    demand = _demand()
    demand["ea_newtons"] = None
    assert _model(demand=demand)["rope_mismatch"] is None


def test_an_unnamed_analysed_rope_is_found_by_its_stiffness():
    demand = _demand()
    demand["net"]["ea_provenance"] = "assumed"
    model = _model(demand=demand, configuration=_configuration(rope="rope-8mm"))
    assert model["rope_mismatch"]["analysed_rope"] == "rope-4mm"


def test_the_route_exports_a_different_rope_and_says_so(client, monkeypatch, tmp_path):
    _plant_demand(monkeypatch, tmp_path)
    folder = tmp_path / "out"
    _into(monkeypatch, folder)
    response = client.post("/api/studies/My Vault/cablenet/exports",
                           json={"configuration": _configuration(rope="rope-8mm")})
    assert response.status_code == 200, response.text
    md = next(Path(p) for p in response.json()["paths"]
              if p.endswith(".md")).read_text(encoding="utf-8")
    assert "not the rope that was analysed" in md


# ---------------------------------------------------------------------------
# The panel writes a force the way the documents do
# ---------------------------------------------------------------------------

# newtons() lives in the model, where the panel's sentences are made; the
# controller beside it writes the rest and is held to the same rule.
_CABLENET_JS = (Path(__file__).resolve().parents[2] / "bench" / "studio"
                / "static" / "cablenet_model.js")
_CABLENET_CONTROLLER_JS = _CABLENET_JS.with_name("cablenet.js")


def test_the_panel_never_rounds_a_force_to_whole_newtons():
    import re
    model = _CABLENET_JS.read_text(encoding="utf-8")
    controller = _CABLENET_CONTROLLER_JS.read_text(encoding="utf-8")
    # a number followed by " N" must come out of newtons(), not toFixed
    for js in (model, controller):
        assert not re.search(r"toFixed\(\d\)\}? N\b", js)
    # the calls, across both files (not the definition, nor kilonewtons)
    assert len(re.findall(r"(?<![A-Za-z])(?<!function )newtons\(", model + controller)) >= 4


def test_the_panels_formatter_renders_what_exports_newtons_renders():
    import json
    import re
    import shutil
    import subprocess
    node = shutil.which("node")
    if node is None:
        pytest.skip("needs node")
    js = _CABLENET_JS.read_text(encoding="utf-8")
    source = re.search(r"function newtons\(value\) \{.*?\n\}", js, re.S).group(0)
    values = [1471.0, 1471, 1470.96, 1470.94, 1500, 1471.4, 0, 450000, 0.04, None]
    out = subprocess.run(
        [node, "-e", source + "console.log(JSON.stringify({}.map(newtons)))".format(
            json.dumps(values))],
        capture_output=True, text=True, check=True).stdout
    assert json.loads(out) == [exports._newtons(v) for v in values]
    # the case that made this more than cosmetic: a fractional ceiling just
    # under the floor must not read as "1471 N against 1471 N"
    assert exports._newtons(1470.94) != exports._newtons(1471.0)


def test_with_eye_m20_chosen_no_rung_drops_to_a_weaker_eye_bolt(
        client, monkeypatch, tmp_path):
    _plant_demand(monkeypatch, tmp_path)
    _into(monkeypatch, tmp_path / "out")
    seen = _spy_on_export_model(monkeypatch)
    chosen = _configuration(chain=["eye-M20", "turnbuckle-eye-eye-M10"])
    response = client.post("/api/studies/My Vault/cablenet/exports",
                           json={"configuration": chosen})
    assert response.status_code == 200, response.text
    assert all(r["configuration"]["chain"][0] == "eye-M20" for r in seen["ladder"])


def test_exactly_one_ladder_row_is_the_chosen_set_even_if_a_rung_equals_it():
    import catalogue
    parts = catalogue.load_parts()
    configuration = _configuration()
    rows = [_scored(parts, configuration), _scored(parts, dict(configuration)),
            _scored(parts, dict(configuration, chain=["eye-M16",
                                                     "turnbuckle-hook-hook-M10"]))]
    model = exports.export_model(parts, _demand(), _row(parts, configuration),
                                 configuration, 10.0, "2026-10-07",
                                 ladder_rows=rows)
    assert [r["is_chosen"] for r in model["ladder"]].count(True) == 1
    assert [r["label"] for r in model["ladder"]].count("Chosen") == 1
    assert len(model["ladder"]) == 2


def test_a_folder_that_cannot_be_listed_is_a_400_naming_it_not_a_500(
        client, monkeypatch, tmp_path):
    _plant_demand(monkeypatch, tmp_path)
    folder = tmp_path / "out"
    _into(monkeypatch, folder)
    real = Path.iterdir

    def refuse(self):
        if self == folder:
            raise PermissionError(13, "Access is denied")
        return real(self)

    monkeypatch.setattr(Path, "iterdir", refuse)
    response = client.post("/api/studies/My Vault/cablenet/exports",
                           json={"configuration": _configuration()})
    assert response.status_code == 400, response.text
    assert str(folder) in response.json()["detail"]
    assert "Access is denied" in response.json()["detail"]


# ---------------------------------------------------------------------------
# The weight, the columns, the sag and the grab: one model, said the same way
# in all three documents
# ---------------------------------------------------------------------------

def _v2_demand(**kwargs):
    demand = _demand()
    demand["schema"] = "bench.cablenet/2"
    # as the engine writes it: the wires are judged at the larger of the entered
    # prestress (300 N) and the fit's greatest wire tension (900 N), here the fit's
    demand["sizing"] = {"stage": "S7", "worst_wire_tension_newtons": 900.0,
                        "fitted_wire_tension_newtons": 900.0, "prestress_newtons": 300.0,
                        "worst_actuator_newtons": 120.0, "worst_sag_mm": 1.4,
                        "load_newtons": 12000.0}
    demand["held"] = {"wire_nodes": [0, 2], "column_heads": [4], "actuators": [1, 3]}
    demand["placement"] = {
        "stage": "S7", "batch": 1, "steps": 10, "reached": True,
        "method": "greedy by unbalanced force, in batches: a heuristic, not an optimum",
        "stranded": [],
        "curve": [{"count": 0, "worst_residual_newtons": 40.0, "worst_sag_mm": 9.0,
                   "residual_norm_newtons": 60.0, "added": []},
                  {"count": 1, "worst_residual_newtons": 20.0, "worst_sag_mm": 4.0,
                   "residual_norm_newtons": 30.0, "added": [1]},
                  {"count": 2, "worst_residual_newtons": 0.0, "worst_sag_mm": 1.4,
                   "residual_norm_newtons": 0.0, "added": [3]}]}
    for stage, sag, column in zip(demand["stages"], (0.5, 1.4), (300.0, 2500.0)):
        stage["node_sag_mm"] = [None, sag, sag / 2, None, None, 0.1, sag / 3, 0.0, 0.2]
        stage["column_forces"] = [{"node": 4, "force": [0.0, 0.0, column],
                                   "newtons": column, "vertical": column}]
        stage["actuator_forces"] = [[0.0, 0.0, 50.0], [0.0, 0.0, 120.0]]
        stage["member_tensions"] = [100.0] * 12
    demand.update(kwargs)
    return demand


def _sized_model(**demand_overrides):
    import catalogue
    demand = _v2_demand(**demand_overrides)
    parts = catalogue.load_parts()
    configuration = _configuration()
    row = _row(parts, configuration,
               load_factor=catalogue.load_factor(parts, configuration, 10.0, demand),
               drive="CL86Y")
    return exports.export_model(parts, demand, row, configuration, 10.0, "2026-10-08")


def test_the_three_documents_say_the_same_load_factor():
    model = _sized_model()
    sentence = exports.load_factor_sentence(model)
    # 1471 N against 900 N: 1.6 passes (1440), 1.7 breaches (1530)
    assert sentence.startswith("Carries 1.6 times the 12.0 kN skin")
    assert "turnbuckle-hook-hook-M10" in sentence
    assert sentence in exports.datasheet_markdown(model)
    hold = exports._sheet_rows(model)["Hold"]
    assert hold[1] == [sentence]
    assert sentence in exports.diagram_svg(model)
    assert model["capacity"]["limit_factor"] == pytest.approx(1.6)


def test_a_load_factor_below_one_leads_with_the_failure_everywhere():
    model = _sized_model(sizing={"stage": "S7", "worst_wire_tension_newtons": 3000.0,
                                 "worst_actuator_newtons": 0.0, "worst_sag_mm": 1.4,
                                 "load_newtons": 12000.0})
    sentence = exports.load_factor_sentence(model)
    assert sentence.startswith("Carries only 0.4 times the 12.0 kN skin, so it does not hold the skin")
    text = exports.datasheet_markdown(model)
    assert text.index(sentence) < text.index("## The load path")
    assert "does not hold" in exports.diagram_svg(model)


def test_without_sizing_every_document_says_not_established_and_names_the_rerun():
    demand = _v2_demand()
    demand.pop("sizing")
    import catalogue
    parts = catalogue.load_parts()
    row = _row(parts, _configuration(), load_factor=None, drive="CL86Y")
    model = exports.export_model(parts, demand, row, _configuration(), 10.0, "2026-10-08")
    sentence = exports.load_factor_sentence(model)
    assert sentence.startswith("Whether it carries the skin is not established")
    assert "run the cable net analysis" in sentence
    assert sentence in exports.datasheet_markdown(model)
    assert exports._sheet_rows(model)["Hold"][1] == [sentence]
    assert sentence in exports.diagram_svg(model)


def test_the_hold_sheet_lists_every_stage_with_its_sag_and_column_force():
    rows = exports._sheet_rows(_sized_model())["Hold"]
    header = rows[3]
    assert header == ["Stage", "Worst sag (mm)", "Nodes past the line",
                      "Worst column force (N)", "At node", "Worst actuator force (N)",
                      "Actuators held"]
    by_name = {row[0]: row for row in rows[4:] if len(row) == 7}
    # cells may be the workbook's float subclasses or the CSV's strings; the
    # number is what is pinned, and the rounding is the data sheet's
    assert float(by_name["S1"][1]) == pytest.approx(0.5) and float(by_name["S7"][1]) == pytest.approx(1.4)
    assert by_name["S7"][2] == 0 and by_name["S1"][2] == 0
    assert float(by_name["S7"][3]) == pytest.approx(2500.0) and by_name["S7"][4] == 4
    assert float(by_name["S7"][5]) == pytest.approx(120.0) and by_name["S7"][6] == 2
    assert rows[2][0].startswith("Grab 2 nodes")


def test_where_to_grab_names_the_count_the_batches_and_the_heuristic():
    model = _sized_model()
    sentence = exports.grab_sentence(model)
    assert sentence.startswith("Grab 2 nodes (2 batches of 1) to bring the net inside the 2.18 mm line")
    assert "heuristic" in sentence
    text = exports.datasheet_markdown(model)
    assert "## Where to grab the net" in text and sentence in text
    assert "nodes 1, 3" in text
    unreached = _sized_model(placement={**_v2_demand()["placement"], "reached": False})
    assert "did not reach the line" in exports.grab_sentence(unreached)


def test_the_data_sheet_has_nine_sections_in_order():
    text = exports.datasheet_markdown(_sized_model())
    headings = [line for line in text.splitlines() if line.startswith("## ")]
    assert headings == [
        "## What this machine replaces", "## What the vault demands", "## What was chosen",
        "## Whether it holds", "## Can it hold the weight", "## Where to grab the net",
        "## The load path", "## The assumptions, listed as assumptions",
        "## What is not checked"]


def test_the_chosen_sheet_gains_the_weight_rows():
    rows = exports._sheet_rows(_sized_model())["Chosen"]
    labels = [row[0] for row in rows if row]
    for label in ("Load factor", "Binds on", "Actuators needed", "Worst column force (N)",
                  "Worst sag (mm)"):
        assert label in labels, label
    load = next(row for row in rows if row and row[0] == "Load factor")
    assert load[1] == "1.6"


def test_the_summaries_read_the_demand_and_compute_nothing_new():
    model = _sized_model()
    assert model["columns"] == {"newtons": 2500.0, "node": 4, "stage": "S7", "vertical": 2500.0}
    assert model["sag"] == {"worst_mm": 1.4, "stage": "S7", "nodes_over_line": 0,
                            "nodes": 6, "acceptance_mm": 2.18}
    assert model["held"]["actuators"] == [1, 3]
    assert model["placement"]["reached"] is True
    bare = exports.export_model(*_sized_model_args_without_v2())
    assert bare["columns"] is None and bare["sag"] is None and bare["placement"] is None


def _sized_model_args_without_v2():
    import catalogue
    parts = catalogue.load_parts()
    return parts, _demand(), _row(parts, _configuration()), _configuration(), 10.0, "2026-10-08"


# ---------------------------------------------------------------------------
# The frames of the raise are shown, and never judged
# ---------------------------------------------------------------------------

def _with_frames(demand):
    """The stages as the engine writes them for a study with a formwork
    document: two frames of the raise first, each with a machine time and no
    course, then every course of the skin, each with a course and no time. The
    net is slack by design in a frame, so its sag is far past the line."""

    frames = []
    for number, time in enumerate((45.0, 100.0), start=1):
        frames.append({
            "stage": number, "name": "F{:g}".format(time), "kind": "raise",
            "time": time, "course": None, "placed_weight_newtons": None,
            "skin_load_sum_newtons": 0.0, "wire_rest_lengths": [2100.0],
            "wire_reel_commands": [0.0], "wire_tensions": [60.0],
            "deviation": 40.0, "reachable": False, "residual_after": 40.0,
            "node_sag_mm": [None, 40.0, 20.0, None, None, 0.1, 13.0, 0.0, 0.2],
            "column_forces": [{"node": 4, "force": [0.0, 0.0, 3000.0],
                               "newtons": 3000.0, "vertical": 3000.0}],
            "actuator_forces": [[0.0, 0.0, 10.0], [0.0, 0.0, 20.0]],
        })
    for index, stage in enumerate(demand["stages"]):
        stage.update(stage=len(frames) + index + 1, time=None, course=index)
    return frames + demand["stages"]


def test_a_frame_instant_of_the_raise_never_enters_the_holds_verdict():
    model = _sized_model(stages=_with_frames(_v2_demand()))
    shape = model["verdict"]["shape"]
    # the frames sit 40 mm past a 2.18 mm line and cannot be reached, and the
    # verdict is the courses', which are inside it
    assert shape["within"] is True
    assert shape["unreachable_stages"] == []
    assert shape["worst_stage"] == "S7" and shape["worst_residual_mm"] == pytest.approx(1.4)
    assert shape["frame_instants"] == ["F45", "F100"]
    assert model["verdict"]["holds"] is True
    text = exports.datasheet_markdown(model)
    section = _section(text, "whether it holds")
    assert ("The 2 frame instants of the raise (F45, F100) are shown below and in the "
            "spreadsheet but are not judged here") in section
    assert "only the courses of the skin are held to the acceptance line" in section
    assert "Stage 1 (F45, raise): residual 40.00 mm, not reachable." in section
    assert "Verdict: holds" in exports.diagram_svg(model)
    # a document with no frames says nothing of them
    assert "frame instant" not in exports.datasheet_markdown(_sized_model())
    assert _sized_model()["verdict"]["shape"]["frame_instants"] == []


def test_one_frame_is_said_in_the_singular():
    stages = _with_frames(_v2_demand())[1:]
    section = _section(exports.datasheet_markdown(_sized_model(stages=stages)),
                       "whether it holds")
    assert ("The frame instant of the raise (F100) is shown below and in the "
            "spreadsheet but is not judged here") in section


def test_a_course_is_judged_though_its_index_is_zero():
    # the first course is course 0: a test of "is there a course" by truth would
    # read it as a frame (here it also carries a time, as a frame does) and let a
    # stage the correction cannot reach go unjudged
    stages = _with_frames(_v2_demand())
    first = next(s for s in stages if s["course"] == 0)
    first.update(reachable=False, time=45.0)
    shape = _sized_model(stages=stages)["verdict"]["shape"]
    assert shape["within"] is False and shape["unreachable_stages"] == [first["name"]]
    # and a course past the line fails whatever the frames do
    stages = _with_frames(_v2_demand())
    stages[-1]["residual_after"] = 7.4
    shape = _sized_model(stages=stages)["verdict"]["shape"]
    assert shape["within"] is False and shape["worst_stage"] == "S7"


def test_a_stage_with_a_course_is_a_course_whatever_its_time():
    # the rule is the ruling's: only an instant with no course is a frame
    stages = _with_frames(_v2_demand())
    last = stages[-1]
    last.update(time=100.0, reachable=False)
    shape = _sized_model(stages=stages)["verdict"]["shape"]
    assert shape["within"] is False and shape["unreachable_stages"] == ["S7"]
    assert shape["frame_instants"] == ["F45", "F100"]


def test_the_frames_are_shown_in_the_hold_sheet_stage_by_stage():
    rows = exports._sheet_rows(_sized_model(stages=_with_frames(_v2_demand())))["Hold"]
    by_name = {row[0]: row for row in rows[4:] if len(row) == 7}
    assert list(by_name) == ["F45", "F100", "S1", "S7"]
    assert float(by_name["F45"][1]) == pytest.approx(40.0)
    assert by_name["F45"][2] == 3 and by_name["S7"][2] == 0
    assert float(by_name["F45"][3]) == pytest.approx(3000.0)


def test_the_sag_summary_judges_the_instants_the_verdict_judges():
    # one worst, not two: the line the verdict reads and the line the grab reads
    # are the same instants, so no sheet says 1.40 mm in one row and 40.00 in the next
    model = _sized_model(stages=_with_frames(_v2_demand()))
    assert model["sag"]["stage"] == "S7"
    assert model["sag"]["worst_mm"] == model["verdict"]["shape"]["worst_residual_mm"]
    assert model["sag"]["nodes_over_line"] == 0
    # a column carries what it carries at every instant, the raise included
    assert model["columns"]["stage"] == "F45" and model["columns"]["newtons"] == 3000.0


# ---------------------------------------------------------------------------
# What a column force and an actuator force are
# ---------------------------------------------------------------------------

def test_the_column_and_actuator_forces_are_said_to_be_one_state_the_fit_chose():
    text = exports.datasheet_markdown(_sized_model())
    for heading in ("hold the weight", "grab the net"):
        section = _section(text, heading)
        assert section.count("one equilibrium state, the one the fit chose") == 1, heading
        assert "not a measurement" in section, heading
        assert "the same for every such state" in section, heading
    columns = _section(text, "hold the weight")
    assert "a member with both ends held carries nothing in the fit" in columns
    assert "The column forces are" in columns
    assert "The actuator forces are" in _section(text, "grab the net")
    # a force that is not reported is not explained
    bare = exports.datasheet_markdown(exports.export_model(*_sized_model_args_without_v2()))
    assert "equilibrium state" not in bare


def test_no_actuator_force_is_explained_when_no_node_is_grabbed():
    model = _sized_model(held={"wire_nodes": [0, 2], "column_heads": [4], "actuators": []})
    text = exports.datasheet_markdown(model)
    assert "equilibrium state" in _section(text, "hold the weight")
    assert "equilibrium state" not in _section(text, "grab the net")


# ---------------------------------------------------------------------------
# Counts agree with their nouns, and what is unknown is blank and not zero
# ---------------------------------------------------------------------------

def test_one_head_is_said_in_the_singular():
    text = exports.datasheet_markdown(_sized_model())
    assert "The columns prop the net at 1 head." in text and "1 heads" not in text
    many = _sized_model(held={"wire_nodes": [0, 2], "column_heads": [4, 5, 6], "actuators": [1, 3]})
    assert "prop the net at 3 heads." in exports.datasheet_markdown(many)


def test_one_node_is_grabbed_in_one_batch_of_up_to_the_size():
    held = {"wire_nodes": [0, 2], "column_heads": [4], "actuators": [1]}
    placement = {**_v2_demand()["placement"], "batch": 20}
    model = _sized_model(held=held, placement=placement)
    sentence = exports.grab_sentence(model)
    assert sentence.startswith(
        "Grab 1 node (1 batch of up to 20) to bring the net inside the 2.18 mm line")
    section = _section(exports.datasheet_markdown(model), "grab the net")
    assert "nodes to grab, in the order the walk chose them: node 1." in section
    assert "nodes 1." not in section


def test_a_last_batch_that_is_not_full_is_not_called_a_full_one():
    held = {"wire_nodes": [0, 2], "column_heads": [4], "actuators": [1, 3, 5]}
    sentence = exports.grab_sentence(
        _sized_model(held=held, placement={**_v2_demand()["placement"], "batch": 2}))
    assert sentence.startswith("Grab 3 nodes (2 batches of up to 2) to bring")
    full = exports.grab_sentence(_sized_model(
        held={**held, "actuators": [1, 3, 5, 7]},
        placement={**_v2_demand()["placement"], "batch": 2}))
    assert full.startswith("Grab 4 nodes (2 batches of 2) to bring")


def test_a_net_that_needs_no_grab_says_so_and_does_not_grab_nothing():
    held = {"wire_nodes": [0, 2], "column_heads": [4], "actuators": []}
    placement = {**_v2_demand()["placement"], "reached": True,
                 "curve": [{"count": 0, "worst_residual_newtons": 0.0, "worst_sag_mm": 1.4,
                            "residual_norm_newtons": 0.0, "added": []}]}
    model = _sized_model(held=held, placement=placement)
    sentence = exports.grab_sentence(model)
    assert sentence.startswith("No node needs grabbing")
    assert "1.40 mm" in sentence and "2.18 mm line" in sentence
    assert "Grab 0" not in sentence and "0 batches" not in sentence
    section = _section(exports.datasheet_markdown(model), "grab the net")
    assert "No node is grabbed." in section
    chosen = {r[0]: r for r in exports._sheet_rows(model)["Chosen"] if r}
    assert chosen["Actuators needed"][1] == 0          # a known zero is a zero


def test_an_unknown_count_of_actuators_is_blank_and_not_zero():
    model = exports.export_model(*_sized_model_args_without_v2())
    rows = exports._sheet_rows(model)
    chosen = {r[0]: r for r in rows["Chosen"] if r}
    assert chosen["Actuators needed"][1] == ""
    assert chosen["Actuators needed"][2].startswith("Where to grab the net is not established")
    assert all(row[6] == "" for row in rows["Hold"][4:] if len(row) == 7)
    text = exports.datasheet_markdown(model)
    assert "No node is grabbed." not in text
    assert "not established" in _section(text, "grab the net")


def test_an_unknown_column_and_an_unknown_sag_read_as_unknown_in_the_chosen_sheet():
    chosen = {r[0]: r for r in exports._sheet_rows(
        exports.export_model(*_sized_model_args_without_v2()))["Chosen"] if r}
    assert chosen["Worst column force (N)"][1] == ""
    assert chosen["Worst column force (N)"][2] == "no column force is recorded"
    assert chosen["Worst sag (mm)"][1] == ""
    assert chosen["Worst sag (mm)"][2] == "no sag is recorded"


def test_nodes_past_the_line_agree_with_their_noun():
    stages = _v2_demand()["stages"]
    stages[1]["node_sag_mm"] = [None, 3.0, 1.0, None, None, 0.1, 2.0, 0.0, 0.2]
    one = {r[0]: r for r in exports._sheet_rows(
        _sized_model(stages=stages))["Chosen"] if r}["Worst sag (mm)"]
    assert one[2] == "at stage S7; 1 node past the line"
    section = _section(exports.datasheet_markdown(_sized_model(stages=stages)), "grab the net")
    assert "and 1 of 6 free nodes is past the line" in section
    stages[1]["node_sag_mm"] = [None, 3.0, 2.5, None, None, 0.1, 2.0, 0.0, 0.2]
    two = {r[0]: r for r in exports._sheet_rows(
        _sized_model(stages=stages))["Chosen"] if r}["Worst sag (mm)"]
    assert two[2] == "at stage S7; 2 nodes past the line"
    section = _section(exports.datasheet_markdown(_sized_model(stages=stages)), "grab the net")
    assert "and 2 of 6 free nodes are past the line" in section
    none = {r[0]: r for r in exports._sheet_rows(_sized_model())["Chosen"] if r}["Worst sag (mm)"]
    assert none[2] == "at stage S7; 0 nodes past the line"


def test_with_no_acceptance_line_no_node_is_counted_against_one():
    model = _sized_model(acceptance=None)
    assert model["sag"]["nodes_over_line"] is None
    section = _section(exports.datasheet_markdown(model), "grab the net")
    assert "no acceptance line is set" in section
    assert "of 6 free nodes" not in section and "None" not in section
    chosen = {r[0]: r for r in exports._sheet_rows(model)["Chosen"] if r}
    assert chosen["Worst sag (mm)"][2] == "at stage S7; no acceptance line is set"
    hold = {r[0]: r for r in exports._sheet_rows(model)["Hold"][4:] if len(r) == 7}
    assert hold["S7"][2] == ""


def test_the_line_is_the_demands_even_when_no_stage_carries_a_sag():
    stages = _v2_demand()["stages"]
    for stage in stages:
        stage.pop("node_sag_mm")
    model = _sized_model(stages=stages)
    assert model["sag"] is None
    sentence = exports.grab_sentence(model)
    assert "the 2.18 mm line" in sentence and "nobody set" not in sentence


# ---------------------------------------------------------------------------
# The load factor sentence for what else can bind
# ---------------------------------------------------------------------------

def test_a_net_that_nothing_binds_is_not_said_to_be_bound_by_nothing():
    light = {"stage": "S7", "worst_wire_tension_newtons": 10.0,
             "fitted_wire_tension_newtons": 4.0, "prestress_newtons": 10.0,
             "worst_actuator_newtons": 0.0, "worst_sag_mm": 1.4, "load_newtons": 12000.0}
    sentence = exports.load_factor_sentence(_sized_model(sizing=light, prestress=10.0))
    assert sentence == ("Carries at least 20.0 times the 12.0 kN skin: nothing binds "
                        "up to that load.")


def test_a_net_that_sags_past_the_line_says_the_shape_binds_with_one_colon():
    sagging = {"stage": "S7", "worst_wire_tension_newtons": 900.0, "worst_actuator_newtons": 0.0,
               "worst_sag_mm": 5.0, "load_newtons": 12000.0}
    model = _sized_model(sizing=sagging)
    assert model["capacity"]["binding_part"] == "shape"
    sentence = exports.load_factor_sentence(model)
    assert sentence == ("Carries only 0.0 times the 12.0 kN skin, so it does not hold "
                        "the skin: the shape binds, because the net sags past the "
                        "acceptance line at any load.")
    assert sentence in exports.datasheet_markdown(model)


def test_a_skin_whose_weight_is_not_recorded_is_not_given_a_weight():
    model = _sized_model()
    model["capacity"] = dict(model["capacity"], skin_newtons=None)
    sentence = exports.load_factor_sentence(model)
    assert sentence.startswith("Carries 1.6 times the skin before") and "kN" not in sentence


def test_a_net_with_no_tension_to_scale_is_not_established_and_says_why():
    # only a document that records no prestress can have nothing to scale: the
    # engine refuses a prestress of zero, and the wires are judged at no less
    still = {"stage": "S7", "worst_wire_tension_newtons": 0.0, "worst_actuator_newtons": 0.0,
             "worst_sag_mm": 1.4, "load_newtons": 12000.0}
    sentence = exports.load_factor_sentence(_sized_model(sizing=still, prestress=None))
    assert sentence == ("Whether it carries the skin is not established: no wire carries "
                        "tension at the sizing stage, so there is nothing to scale.")


# ---------------------------------------------------------------------------
# One figure, one way to write it, in every document
# ---------------------------------------------------------------------------

def test_the_worst_column_force_reads_the_same_in_all_three_documents():
    model = _sized_model()
    shown = exports._newtons(model["columns"]["newtons"])
    assert shown == "2500.0"
    rows = exports._sheet_rows(model)
    chosen = next(r for r in rows["Chosen"] if r and r[0] == "Worst column force (N)")
    hold = {r[0]: r for r in rows["Hold"][4:] if len(r) == 7}["S7"]
    assert exports._newtons(chosen[1]) == shown == exports._newtons(hold[3])
    assert "{} N at node 4 at stage S7".format(shown) in exports.datasheet_markdown(model)
    assert chosen[2] == "at node 4 at stage S7"


def test_the_new_sections_write_every_force_through_the_one_formatter():
    import re
    model = _sized_model()
    for text in (exports.datasheet_markdown(model), exports.diagram_svg(model)):
        assert not re.search(r"\b\d{4,} N\b", text)
        assert _grouped_figures(text) == [], _grouped_figures(text)
    md = exports.datasheet_markdown(model)
    assert "2500.0 N" in md and "120.0 N" in md and "12.0 kN" in md
    assert "\u2014" not in md and "$" not in md


def test_the_hold_sheet_survives_the_csv_fallback(tmp_path, monkeypatch):
    import csv
    monkeypatch.setattr(exports, "_openpyxl", None)
    model = _sized_model()
    written = exports.write_spreadsheet(model, tmp_path, "x")
    path = next(p for p in written if p.name.endswith("-hold.csv"))
    with open(path, newline="", encoding="utf-8-sig") as handle:
        rows = list(csv.reader(handle))
    assert rows[1] == [exports.load_factor_sentence(model)]
    assert rows[2] == [exports.grab_sentence(model)]
    s7 = next(r for r in rows if r[:1] == ["S7"])
    assert s7 == ["S7", "1.40", "0", "2500.0", "4", "120.0", "2"]


def test_the_hold_sheet_carries_forces_and_lengths_as_numbers(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    written = exports.write_spreadsheet(_sized_model(), tmp_path, "x")
    sheet = openpyxl.load_workbook(written[0])["Hold"]
    row = next(r for r in sheet.iter_rows() if r[0].value == "S7")
    assert row[3].value == 2500.0 and row[3].number_format == "0.0"
    assert row[1].value == pytest.approx(1.4) and row[1].number_format == "0.00"
    assert sheet.cell(row=1, column=1).font.bold


def test_the_diagram_is_wide_enough_for_the_line_it_adds():
    import re
    model = _sized_model()
    model["capacity"] = {"limit_factor": None, "detail": "x" * 300}
    sentence = exports.load_factor_sentence(model)
    svg = exports.diagram_svg(model)
    width = int(re.search(r'<svg[^>]* width="(\d+)"', svg).group(1))
    # the width is rounded to a whole pixel, as it is for the heading and the warning
    assert exports._estimated_width(sentence, 12) + 32 <= width + 0.5
    assert 'viewBox="0 0 {} '.format(width) in svg
    assert sentence in svg


def test_the_load_factor_line_sits_under_the_verdict_and_above_the_rope_warning():
    import re
    import catalogue
    parts = catalogue.load_parts()
    configuration = _configuration(rope="rope-8mm")
    demand = _v2_demand()
    row = _row(parts, configuration, drive="CL86Y",
               load_factor=catalogue.load_factor(parts, configuration, 10.0, demand))
    model = exports.export_model(parts, demand, row, configuration, 10.0, "2026-10-08")
    assert model["rope_mismatch"]
    svg = exports.diagram_svg(model)
    placed = {text: float(y) for y, text in re.findall(
        r'<text x="[^"]*" y="([^"]*)"[^>]*>([^<]*)</text>', svg)}
    verdict = next(y for text, y in placed.items() if text.startswith("Verdict:"))
    warning = next(y for text, y in placed.items() if text.startswith("Warning:"))
    assert verdict < placed[exports.load_factor_sentence(model)] < warning
    height = float(re.search(r'<svg[^>]* height="(\d+)"', svg).group(1))
    assert max(placed.values()) < height


def test_the_chosen_sheet_names_what_binds_and_leaves_blank_what_is_unknown():
    def row(model, label):
        return next(r for r in exports._sheet_rows(model)["Chosen"] if r and r[0] == label)

    assert row(_sized_model(), "Binds on")[1] == "turnbuckle-hook-hook-M10"
    sagging = {"stage": "S7", "worst_wire_tension_newtons": 900.0, "worst_actuator_newtons": 0.0,
               "worst_sag_mm": 5.0, "load_newtons": 12000.0}
    assert row(_sized_model(sizing=sagging), "Binds on")[1] == "shape"
    unsized = exports.export_model(*_sized_model_args_without_v2())
    assert row(unsized, "Load factor")[1] == "" and row(unsized, "Binds on")[1] == ""
    assert row(unsized, "Binds on")[2] == ""


def test_an_unsized_diagram_keeps_its_verdict_and_adds_the_line_that_says_so():
    # what the diagram said of a document from before the sizing block, unchanged,
    # and then the one thing it could not say before
    passing = exports.diagram_svg(_model())
    assert "Verdict: holds" in passing
    assert "Whether it carries the skin is not established" in passing
    demand = _demand()
    demand["stages"][1]["residual_after"] = 7.4
    failing = exports.diagram_svg(_model(demand=demand))
    assert "Verdict: does not hold" in failing
    assert "Whether it carries the skin is not established" in failing


def test_the_term_to_part_map_is_the_catalogues_alone():
    assert not hasattr(exports, "_term_part")


# ---------------------------------------------------------------------------
# A document that lacks a block says so, in every renderer, and nothing throws
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("dropped", [
    (), ("sizing",), ("held",), ("placement",), ("sizing", "held"),
    ("sizing", "placement"), ("held", "placement"), ("sizing", "held", "placement")])
@pytest.mark.parametrize("bare_stages", [False, True])
def test_a_demand_that_lacks_a_block_never_throws_and_never_prints_none(
        dropped, bare_stages, tmp_path):
    import re
    import catalogue
    parts = catalogue.load_parts()
    configuration = _configuration()
    demand = _v2_demand()
    for key in dropped:
        demand.pop(key)
    if bare_stages:
        for stage in demand["stages"]:
            for key in ("node_sag_mm", "column_forces", "actuator_forces"):
                stage.pop(key)
    row = _row(parts, configuration, drive="CL86Y",
               load_factor=catalogue.load_factor(parts, configuration, 10.0, demand))
    model = exports.export_model(parts, demand, row, configuration, 10.0, "2026-10-08")
    exports.write_spreadsheet(model, tmp_path, "x")
    exports.write_diagram(model, tmp_path, "x")
    exports.write_datasheet(model, tmp_path, "x")
    sheets = exports._sheet_rows(model)
    for name in ("Chosen", "Hold"):
        flat = " ".join(str(cell) for row in sheets[name] for cell in row)
        assert "None" not in flat, (name, dropped, bare_stages)
    text = exports.datasheet_markdown(model)
    assert "None" not in text and not re.search(r"\bnan\b", text.lower()), dropped
    if "sizing" in dropped:
        assert "Whether it carries the skin is not established" in text
    if "placement" in dropped:
        assert "Where to grab the net is not established" in text
    if "held" in dropped and "placement" not in dropped:
        # the placement is there and the held set is not: an empty set is not assumed
        assert ("Where to grab the net is not established: the demand document has "
                "no held set") in text
        for sentence in ("No node needs grabbing", "With none grabbed", "With them held",
                         "0 heads"):
            assert sentence not in text, (sentence, dropped, bare_stages)
        chosen = {r[0]: r for r in sheets["Chosen"] if r}
        assert chosen["Actuators needed"][1] == ""
        assert "no held set" in chosen["Actuators needed"][2]


def test_a_column_with_no_force_in_it_leaves_its_cells_blank():
    # a column with no force is not the worst column of its stage, in the sheet
    # as in the summary, so the sheet names no node for it either
    stages = _v2_demand()["stages"]
    stages[1]["column_forces"] = [{"node": 4}]
    model = _sized_model(stages=stages)
    hold = {r[0]: r for r in exports._sheet_rows(model)["Hold"][4:] if len(r) == 7}
    assert hold["S7"][3] == "" and hold["S7"][4] == ""
    assert model["columns"]["stage"] == "S1"


# ---------------------------------------------------------------------------
# The route: one request, three documents, one sentence
# ---------------------------------------------------------------------------

def test_the_route_writes_the_weight_and_the_grab_into_all_three(client, monkeypatch, tmp_path):
    _plant_demand(monkeypatch, tmp_path, demand=_v2_demand())
    folder = tmp_path / "out"
    _into(monkeypatch, folder)
    response = client.post("/api/studies/My Vault/cablenet/exports",
                           json={"configuration": _configuration()})
    assert response.status_code == 200, response.text
    written = [Path(p) for p in response.json()["paths"]]
    md = next(p for p in written if p.suffix == ".md").read_text(encoding="utf-8")
    svg = next(p for p in written if p.suffix == ".svg").read_text(encoding="utf-8")
    weight = "Carries 1.6 times the 12.0 kN skin before turnbuckle-hook-hook-M10 binds."
    grab = "Grab 2 nodes (2 batches of 1) to bring the net inside the 2.18 mm line"
    assert weight in md and weight in svg
    assert grab in md
    assert "## Can it hold the weight" in md and "## Where to grab the net" in md
    if exports._openpyxl is not None:
        import openpyxl
        book = openpyxl.load_workbook(next(p for p in written if p.suffix == ".xlsx"))
        assert book.sheetnames[-1] == "Hold"
        hold = [[c.value for c in row] for row in book["Hold"].iter_rows()]
        assert hold[1][0] == weight and hold[2][0].startswith(grab)


# ---------------------------------------------------------------------------
# A held set that is not there is not an empty one
# ---------------------------------------------------------------------------

def test_a_missing_held_set_is_not_read_as_an_empty_one():
    # a placement, a curve whose first figure is 9.00 mm against a 2.18 mm line,
    # column forces, and no held set: nothing here says that no node is grabbed
    model = _sized_model(held=None)
    sentence = exports.grab_sentence(model)
    assert sentence == ("Where to grab the net is not established: the demand document "
                        "has no held set; run the cable net analysis again.")
    text = exports.datasheet_markdown(model)
    assert sentence in _section(text, "grab the net")
    assert "No node needs grabbing" not in text and "already inside" not in text
    assert "With none grabbed" not in text and "With them held" not in text
    assert "0 heads" not in text
    weight = _section(text, "hold the weight")
    assert ("How many heads the columns prop is not recorded, because the demand "
            "document's held set does not list them.") in weight
    assert "The worst column force is 2500.0 N at node 4 at stage S7" in weight
    sheets = exports._sheet_rows(model)
    chosen = {r[0]: r for r in sheets["Chosen"] if r}
    assert chosen["Actuators needed"][1] == "" and chosen["Actuators needed"][2] == sentence
    assert sheets["Hold"][2] == [sentence]
    assert all(row[6] == "" for row in sheets["Hold"][4:] if len(row) == 7)
    # what the document does record is still said, without a claim about the held set
    grab = _section(text, "grab the net")
    assert "The worst first-order sag is 1.40 mm at stage S7, and 0 of 6 free nodes are past the line." in grab
    assert "equilibrium state" not in grab


@pytest.mark.parametrize("held", [
    None, {}, {"wire_nodes": [0, 2], "column_heads": [4]}, {"actuators": None},
    {"actuators": "abc"}, {"actuators": 3}])
def test_an_actuator_list_that_is_not_a_list_is_not_established(held):
    model = _sized_model(held=held)
    sentence = exports.grab_sentence(model)
    assert sentence.startswith("Where to grab the net is not established")
    assert "no held set" in sentence
    chosen = {r[0]: r for r in exports._sheet_rows(model)["Chosen"] if r}
    assert chosen["Actuators needed"][1] == ""
    assert "No node needs grabbing" not in exports.datasheet_markdown(model)


@pytest.mark.parametrize("heads", [None, "abc", 4])
def test_a_held_set_with_no_list_of_heads_does_not_count_them(heads):
    held = {"wire_nodes": [0, 2], "column_heads": heads, "actuators": [1, 3]}
    text = exports.datasheet_markdown(_sized_model(held=held))
    weight = _section(text, "hold the weight")
    assert "does not list them" in weight and "0 heads" not in text
    assert "The worst column force is 2500.0 N at node 4 at stage S7" in weight
    # the actuators are known, so the grab says what it said before
    grab = _section(text, "grab the net")
    assert "With them held, the worst first-order sag is 1.40 mm" in grab
    # and a known empty list is a count of none, not an unknown
    none = _section(exports.datasheet_markdown(_sized_model(
        held={"wire_nodes": [0, 2], "column_heads": [], "actuators": [1, 3]})), "hold the weight")
    assert "The columns prop the net at 0 heads." in none


# ---------------------------------------------------------------------------
# One reader per stage: the summaries and the Hold sheet cannot disagree
# ---------------------------------------------------------------------------

def test_the_model_holds_one_entry_per_stage_read_once():
    model = _sized_model(stages=_with_frames(_v2_demand()))
    hold = model["hold"]
    assert [entry["name"] for entry in hold] == ["F45", "F100", "S1", "S7"]
    assert [entry["frame"] for entry in hold] == [True, True, False, False]
    s7 = hold[-1]
    assert s7["worst_sag_mm"] == pytest.approx(1.4) and s7["nodes"] == 6
    assert s7["nodes_over_line"] == 0
    assert s7["worst_column"] == {"node": 4, "newtons": 2500.0, "vertical": 2500.0}
    assert s7["worst_actuator_newtons"] == pytest.approx(120.0)
    assert hold[0]["nodes_over_line"] == 3 and hold[0]["worst_column"]["newtons"] == 3000.0
    # a v1 stage has nothing to read, and every figure says so
    bare = exports.export_model(*_sized_model_args_without_v2())["hold"]
    assert [entry["name"] for entry in bare] == ["S1", "S7"]
    assert all(entry["worst_sag_mm"] is None and entry["nodes_over_line"] is None
               and entry["worst_column"] is None and entry["worst_actuator_newtons"] is None
               for entry in bare)


def test_the_chosen_and_hold_sheets_read_one_figure_per_stage():
    stages = _with_frames(_v2_demand())
    course = next(s for s in stages if s["name"] == "S7")
    course["node_sag_mm"] = [None, 3.0, 2.5, None, None, 0.1, 2.0, 0.0, 0.2]   # two past 2.18
    model = _sized_model(stages=stages)
    rows = exports._sheet_rows(model)
    hold = {r[0]: r for r in rows["Hold"][4:] if len(r) == 7}
    chosen = {r[0]: r for r in rows["Chosen"] if r}
    # the same cell and the same count for the stage the summary names
    assert model["sag"]["stage"] == "S7"
    assert hold["S7"][1] == chosen["Worst sag (mm)"][1] == pytest.approx(3.0)
    assert hold["S7"][2] == 2
    assert chosen["Worst sag (mm)"][2] == "at stage S7; {} past the line".format(
        exports._count(hold["S7"][2], "node"))
    # and the column the summary names is the column the sheet names for that stage
    named = model["columns"]
    assert chosen["Worst column force (N)"][1] == hold[named["stage"]][3]
    assert chosen["Worst column force (N)"][2] == "at node {} at stage {}".format(
        hold[named["stage"]][4], named["stage"])


def test_the_summaries_and_the_hold_sheet_read_the_one_stage_reader(monkeypatch):
    def doctored(stage, acceptance):
        return {"name": stage.get("name"), "frame": False, "nodes": 4, "nodes_over_line": 3,
                "worst_sag_mm": 11.0 if stage.get("name") == "S7" else 5.0,
                "worst_column": {"node": 9, "newtons": 777.0, "vertical": 700.0},
                "worst_actuator_newtons": 66.0}

    monkeypatch.setattr(exports, "_stage_hold", doctored)
    model = _sized_model()
    assert model["sag"] == {"worst_mm": 11.0, "stage": "S7", "nodes_over_line": 3,
                            "nodes": 4, "acceptance_mm": 2.18}
    assert model["columns"] == {"newtons": 777.0, "node": 9, "stage": "S1", "vertical": 700.0}
    hold = {r[0]: r for r in exports._sheet_rows(model)["Hold"][4:] if len(r) == 7}
    assert float(hold["S7"][1]) == 11.0 and hold["S7"][2] == 3
    assert float(hold["S7"][3]) == 777.0 and hold["S7"][4] == 9 and float(hold["S7"][5]) == 66.0


def test_the_hold_sheet_prints_what_the_model_holds_and_computes_nothing():
    model = _sized_model()
    for entry in model["hold"]:
        entry.update(worst_sag_mm=123.0, nodes_over_line=7, worst_actuator_newtons=45.0,
                     worst_column={"node": 8, "newtons": 999.0, "vertical": 1.0})
    hold = {r[0]: r for r in exports._sheet_rows(model)["Hold"][4:] if len(r) == 7}
    for name in ("S1", "S7"):
        assert [float(hold[name][1]), hold[name][2], float(hold[name][3]), hold[name][4],
                float(hold[name][5])] == [123.0, 7, 999.0, 8, 45.0]


# ---------------------------------------------------------------------------
# A value that is not a number is skipped, by every reader, and never raised on
# ---------------------------------------------------------------------------

def test_a_column_force_that_is_not_a_number_is_skipped_and_not_raised(tmp_path):
    stages = _v2_demand()["stages"]
    stages[1]["column_forces"] = [
        {"node": 4, "force": [0.0, 0.0, 1.0], "newtons": "abc", "vertical": 1.0},
        {"node": 6, "force": [0.0, 0.0, 1.0], "newtons": True, "vertical": 1.0},
        {"node": 7, "force": [0.0, 0.0, 1.0], "newtons": float("nan"), "vertical": 1.0},
        {"node": 8, "force": [0.0, 0.0, 1.0], "newtons": None, "vertical": 1.0},
        {"node": 5, "force": [0.0, 0.0, 700.0], "newtons": 700.0, "vertical": 700.0}]
    model = _sized_model(stages=stages)
    # every document is written, and the summary and the sheet name the same column
    exports.write_spreadsheet(model, tmp_path, "x")
    exports.write_diagram(model, tmp_path, "x")
    exports.write_datasheet(model, tmp_path, "x")
    assert model["columns"] == {"newtons": 700.0, "node": 5, "stage": "S7", "vertical": 700.0}
    hold = {r[0]: r for r in exports._sheet_rows(model)["Hold"][4:] if len(r) == 7}
    assert float(hold["S7"][3]) == 700.0 and hold["S7"][4] == 5
    # a stage whose columns are all junk leaves its cells blank and the summary alone
    stages[1]["column_forces"] = [{"node": 4, "newtons": "abc"}, {"node": 6, "newtons": [1]}]
    model = _sized_model(stages=stages)
    exports.write_spreadsheet(model, tmp_path, "y")
    hold = {r[0]: r for r in exports._sheet_rows(model)["Hold"][4:] if len(r) == 7}
    assert hold["S7"][3] == "" and hold["S7"][4] == ""
    assert model["columns"]["stage"] == "S1" and model["columns"]["newtons"] == 300.0


def test_a_sag_or_an_actuator_force_that_is_not_a_number_is_skipped(tmp_path):
    stages = _v2_demand()["stages"]
    stages[1]["node_sag_mm"] = [None, "x", True, float("nan"), float("inf"), 1.0, 3.0]
    stages[1]["actuator_forces"] = [[0.0, 0.0, "x"], [3.0, 4.0, 0.0], None, "abc", [1.0, None, 2.0]]
    model = _sized_model(stages=stages)
    exports.write_spreadsheet(model, tmp_path, "x")
    s7 = {entry["name"]: entry for entry in model["hold"]}["S7"]
    # only 1.0 and 3.0 are numbers: two free nodes, one of them past the 2.18 line
    assert s7["worst_sag_mm"] == 3.0 and s7["nodes"] == 2 and s7["nodes_over_line"] == 1
    assert s7["worst_actuator_newtons"] == 5.0
    assert model["sag"]["worst_mm"] == 3.0 and model["sag"]["nodes"] == 2
    hold = {r[0]: r for r in exports._sheet_rows(model)["Hold"][4:] if len(r) == 7}
    assert float(hold["S7"][1]) == 3.0 and hold["S7"][2] == 1 and float(hold["S7"][5]) == 5.0
    # a stage that carries only junk has nothing to say, and says it with blanks
    stages[1]["node_sag_mm"] = ["x", None, True]
    stages[1]["actuator_forces"] = ["abc"]
    model = _sized_model(stages=stages)
    hold = {r[0]: r for r in exports._sheet_rows(model)["Hold"][4:] if len(r) == 7}
    assert hold["S7"][1] == "" and hold["S7"][2] == "" and hold["S7"][5] == ""
    assert model["sag"]["stage"] == "S1"


def test_a_stage_whose_arrays_are_not_lists_is_read_as_having_none():
    stages = _v2_demand()["stages"]
    stages[1].update(node_sag_mm="abc", column_forces=5, actuator_forces={"a": 1})
    model = _sized_model(stages=stages)
    entry = {e["name"]: e for e in model["hold"]}["S7"]
    assert entry["worst_sag_mm"] is None and entry["worst_column"] is None
    assert entry["worst_actuator_newtons"] is None
    exports._sheet_rows(model)
    exports.datasheet_markdown(model)


def test_the_worst_column_is_found_by_its_force_and_named_by_its_own_node():
    # the engine's order is [9, 4, 7] and held.column_heads is [4, 7, 9]: the worst
    # is the one with the force, and it is named by the node its own entry carries
    stages = _v2_demand()["stages"]
    stages[1]["column_forces"] = [
        {"node": 9, "force": [0.0, 0.0, 100.0], "newtons": 100.0, "vertical": 100.0},
        {"node": 4, "force": [0.0, 0.0, 2500.0], "newtons": 2500.0, "vertical": 2400.0},
        {"node": 7, "force": [0.0, 0.0, 900.0], "newtons": 900.0, "vertical": 900.0}]
    held = {"wire_nodes": [0, 2], "column_heads": [4, 7, 9], "actuators": [1, 3]}
    model = _sized_model(stages=stages, held=held)
    assert model["columns"] == {"newtons": 2500.0, "node": 4, "stage": "S7", "vertical": 2400.0}
    hold = {r[0]: r for r in exports._sheet_rows(model)["Hold"][4:] if len(r) == 7}
    assert float(hold["S7"][3]) == 2500.0 and hold["S7"][4] == 4
    assert ("The columns prop the net at 3 heads. The worst column force is 2500.0 N at "
            "node 4 at stage S7, of which 2400.0 N is vertical.") in exports.datasheet_markdown(model)


def test_a_column_entry_that_is_not_an_entry_is_skipped():
    stages = _v2_demand()["stages"]
    stages[1]["column_forces"] = [5, "abc", None, [1, 2],
                                  {"node": 5, "force": [0.0, 0.0, 700.0],
                                   "newtons": 700.0, "vertical": 700.0}]
    model = _sized_model(stages=stages)
    assert model["columns"]["newtons"] == 700.0 and model["columns"]["node"] == 5
    exports._sheet_rows(model)


# ---------------------------------------------------------------------------
# The floor every wire is judged at: the larger of the entered prestress and the
# greatest tension the fit found, where it is set, and the stage that sizes the
# parts as another clause
# ---------------------------------------------------------------------------

def _fit_wins(figure):
    """The sizing block as the engine writes it when the fit's greatest wire
    tension, `figure`, is above the 300 N the study entered."""

    return {**_v2_demand()["sizing"], "worst_wire_tension_newtons": figure,
            "fitted_wire_tension_newtons": figure}


def test_the_floor_names_the_instant_where_the_largest_wire_tension_occurs():
    # the raise carries the largest wire tension in the document, and the stage
    # that sizes the parts is a course
    stages = _with_frames(_v2_demand())
    stages[0]["wire_tensions"] = [1500.0]            # F45, a frame of the raise
    model = _sized_model(stages=stages, sizing=_fit_wins(1500.0))
    assert model["demand"]["prestress_floor_newtons"] == 1500.0
    assert model["demand"]["prestress_floor_instant"] == {"name": "F45", "frame": True}
    section = _section(exports.datasheet_markdown(model), "demands")
    assert ("The net is held at a prestress floor of 1500.0 N, the larger of the entered "
            "prestress (300.0 N) and the greatest tension the fit found in any wire "
            "(1500.0 N), reached at the raise's instant F45. The stage that sizes the "
            "parts is S7.") in section
    assert "sizes it" not in section


def test_a_floor_reached_in_a_course_reads_as_a_stage():
    section = _section(exports.datasheet_markdown(_sized_model()), "demands")
    assert ("the greatest tension the fit found in any wire (900.0 N), reached at stage "
            "S7. The stage that sizes the parts is S7.") in section
    # and the stage that sizes the parts can be another course than the one that
    # reaches the floor
    stages = _v2_demand()["stages"]
    stages[0]["wire_tensions"] = [1000.0]            # S1 carries more than S7's 900.0
    model = _sized_model(stages=stages, sizing=_fit_wins(1000.0))
    assert model["demand"]["prestress_floor_instant"] == {"name": "S1", "frame": False}
    section = _section(exports.datasheet_markdown(model), "demands")
    assert "reached at stage S1. The stage that sizes the parts is S7." in section


def test_a_floor_two_instants_share_is_named_at_the_first_in_the_documents_order():
    stages = _with_frames(_v2_demand())
    stages[0]["wire_tensions"] = [1500.0]
    stages[1]["wire_tensions"] = [1500.0]
    model = _sized_model(stages=stages, sizing=_fit_wins(1500.0))
    assert model["demand"]["prestress_floor_instant"] == {"name": "F45", "frame": True}


def test_the_instant_is_the_one_whose_own_worst_tension_is_the_floor():
    # several wires: the stage is found by its worst wire, not by its first
    stages = _v2_demand()["stages"]
    stages[0]["wire_tensions"] = [100.0, 1200.0, 50.0]
    stages[1]["wire_tensions"] = [900.0, 1100.0]
    model = _sized_model(stages=stages, sizing=_fit_wins(1200.0))
    assert model["demand"]["prestress_floor_newtons"] == 1200.0
    assert model["demand"]["prestress_floor_instant"] == {"name": "S1", "frame": False}


def test_a_document_with_no_wire_tension_names_no_instant_for_the_floor():
    # no sizing block and no wire tension: the entered prestress is the floor, set by
    # it at every instant, and with no prestress either there is no floor at all
    for stages in ([], [{"stage": 1, "name": "S1", "wire_tensions": []}]):
        model = _sized_model(stages=stages, sizing=None)
        assert model["demand"]["prestress_floor_newtons"] == 300.0
        assert model["demand"]["prestress_floor_instant"] == {"prestress": True}
        section = _section(exports.datasheet_markdown(model), "demands")
        assert "reached at" not in section and "The stage that sizes the parts is S7." in section
        assert ("the greatest tension the fit found in any wire (not recorded), set by the "
                "entered prestress.") in section
        bare = _sized_model(stages=stages, sizing=None, prestress=None)
        assert bare["demand"]["prestress_floor_newtons"] is None
        assert bare["demand"]["prestress_floor_instant"] is None
        section = _section(exports.datasheet_markdown(bare), "demands")
        assert ("No prestress floor is recorded: the document gives neither an entered "
                "prestress nor a wire tension. The stage that sizes the parts is S7.") in section
        assert "not recorded N" not in section.split("\n\n")[1]


def test_the_other_two_documents_still_name_the_stage_that_sizes_the_parts():
    stages = _with_frames(_v2_demand())
    stages[0]["wire_tensions"] = [1500.0]
    model = _sized_model(stages=stages, sizing=_fit_wins(1500.0))
    assert "sized at stage S7" in exports.diagram_svg(model)
    chosen = {r[0]: r for r in exports._sheet_rows(model)["Chosen"] if r}
    assert float(chosen["Prestress the build demands (N)"][1]) == 1500.0


def _held_at_the_prestress(prestress=3000.0, fitted=14.6):
    """The sizing block the engine writes when the fit's wires carry less than the
    prestress entered, which is the real study's case (14.6 N at 300 N)."""

    return {**_v2_demand()["sizing"], "worst_wire_tension_newtons": prestress,
            "fitted_wire_tension_newtons": fitted, "prestress_newtons": prestress}


def test_the_wires_are_judged_at_the_entered_prestress_when_the_fit_finds_less():
    # a re-run at 3000 N of a net whose fit carries 14.6 N in its worst wire: the
    # turnbuckle's 1471 N does not carry it, however little the fit finds
    import catalogue
    parts = catalogue.load_parts()
    demand = _v2_demand(prestress=3000.0, sizing=_held_at_the_prestress())
    floor = catalogue.wire_floor(demand)
    assert floor == {"newtons": 3000.0, "fitted_newtons": 14.6, "prestress_newtons": 3000.0}
    row = _row(parts, _configuration(), passes=1471.0 >= floor["newtons"],
               margin=1471.0 / floor["newtons"],
               load_factor=catalogue.load_factor(parts, _configuration(), 10.0, demand))
    model = exports.export_model(parts, demand, row, _configuration(), 10.0, "2026-10-08")
    assert model["demand"]["prestress_floor_newtons"] == 3000.0
    assert model["verdict"]["tension"]["prestress_floor_newtons"] == 3000.0
    assert model["demand"]["prestress_floor_instant"] == {"prestress": True}
    assert model["capacity"]["worst_wire_tension_newtons"] == 3000.0
    assert model["capacity"]["sufficient"] is False
    assert model["verdict"]["holds"] is False
    text = exports.datasheet_markdown(model)
    section = _section(text, "demands")
    assert ("The net is held at a prestress floor of 3000.0 N, the larger of the entered "
            "prestress (3000.0 N) and the greatest tension the fit found in any wire "
            "(14.6 N), set by the entered prestress. The stage that sizes the parts is S7. "
            "The prestress is the floor every member is given; the sag is judged at it."
            ) in section
    assert "against a prestress floor of 3000.0 N" in _section(text, "whether it holds")
    assert exports.load_factor_sentence(model).startswith("Carries only 0.4 times")
    chosen = {r[0]: r for r in exports._sheet_rows(model)["Chosen"] if r}
    assert float(chosen["Prestress the build demands (N)"][1]) == 3000.0


def test_a_document_from_before_the_two_figures_is_judged_by_the_same_rule():
    # the real study's block carries the fit's figure alone: 14.6 N at 300 N entered
    older = {"stage": "S17", "worst_wire_tension_newtons": 14.6,
             "worst_actuator_newtons": 435.1, "worst_sag_mm": 1.4, "load_newtons": 12000.0}
    model = _sized_model(sizing=older)
    assert model["demand"]["prestress_floor_newtons"] == 300.0
    assert model["demand"]["fitted_wire_tension_newtons"] == 14.6
    assert model["demand"]["entered_prestress_newtons"] == 300.0
    assert model["capacity"]["worst_wire_tension_newtons"] == 300.0
    assert ("a prestress floor of 300.0 N, the larger of the entered prestress (300.0 N) "
            "and the greatest tension the fit found in any wire (14.6 N), set by the "
            "entered prestress.") in exports.datasheet_markdown(model)
    # and the fit's figure stands when it is the larger
    assert _sized_model(sizing={**older, "worst_wire_tension_newtons": 900.0})[
        "demand"]["prestress_floor_newtons"] == 900.0


def test_a_first_version_document_says_the_floor_rule_in_its_own_words():
    # no sizing block: the wires are read stage by stage (900 N at S7) against the
    # prestress the study entered, and the cut rule is still the cut rule
    model = _model()
    assert model["demand"]["prestress_floor_newtons"] == 900.0
    section = _section(exports.datasheet_markdown(model), "demands")
    assert ("The net is held at a prestress floor of 900.0 N, the larger of the entered "
            "prestress (300.0 N) and the largest tension any wire carries at any stage of "
            "the build (900.0 N), reached at stage S7. The stage that sizes the parts is S7. "
            "The study entered a uniform prestress of 300.0 N for the cut rule; the staged "
            "analysis then found what the wires actually have to carry.") in section
    assert "the fit" not in section
    heavy = _model(demand=_demand(prestress=1000.0))
    assert heavy["demand"]["prestress_floor_newtons"] == 1000.0
    assert ("(900.0 N), set by the entered prestress. The stage that sizes the parts is S7."
            in _section(exports.datasheet_markdown(heavy), "demands"))


# ---------------------------------------------------------------------------
# A version 2 document is described in its own words: the net is held by its
# wires, its column heads and its grabbed nodes, and nothing is corrected or cut.
# The first version's words are untouched.
# ---------------------------------------------------------------------------

def _chosen_row(model, label):
    return next(r for r in exports._sheet_rows(model)["Chosen"] if r and r[0] == label)


def test_the_shape_half_of_a_version_2_document_is_the_sag_with_the_grabbed_nodes_held():
    section = _section(exports.datasheet_markdown(_sized_model()), "whether it holds")
    assert ("The shape half asks whether the net, with the grabbed nodes held, stays within "
            "the acceptance line. The worst sag with the grabbed nodes held is 1.40 mm, at "
            "stage S7, against an acceptance line of 2.18 mm, which is within it.") in section
    assert "once corrected" not in section and "after correction" not in section
    stages = _v2_demand()["stages"]
    stages[1].update(residual_after=3.0, reachable=False)
    section = _section(exports.datasheet_markdown(_sized_model(stages=stages)),
                       "whether it holds")
    assert ("which is outside it. The grabbed nodes do not bring stage S7 inside the line, "
            "which fails the half outright.") in section
    assert "The correction does not reach" not in section
    unknown = _section(exports.datasheet_markdown(_sized_model(acceptance=None)),
                       "whether it holds")
    assert ("The shape half asks whether the net, with the grabbed nodes held, stays within "
            "the acceptance line. That has not been established, because no acceptance "
            "line is set") in unknown


def test_the_first_versions_shape_half_keeps_its_own_words():
    model = _model()
    model["verdict"]["shape"]["unreachable_stages"] = ["S7"]
    section = _section(exports.datasheet_markdown(model), "whether it holds")
    assert ("The shape half asks whether the net, once corrected, stays within the "
            "acceptance line. The worst residual after correction is 1.40 mm") in section
    assert "The correction does not reach stage S7, which fails the half outright." in section
    assert "grabbed" not in section


def test_a_version_2_net_is_held_by_its_wires_its_column_heads_and_its_grabbed_nodes():
    section = _section(exports.datasheet_markdown(_sized_model()), "demands")
    assert "The net is held by 2 wires, 1 column head and 2 grabbed nodes." in section
    assert "anchors" not in section and "for the cut rule" not in section
    # the real study's counts, from its held set
    held = {"wire_nodes": list(range(105)), "column_heads": list(range(46)),
            "actuators": list(range(800))}
    section = _section(exports.datasheet_markdown(_sized_model(held=held)), "demands")
    assert "The net is held by 105 wires, 46 column heads and 800 grabbed nodes." in section
    # a list the set does not carry is not counted as none
    section = _section(exports.datasheet_markdown(
        _sized_model(held={"wire_nodes": [0, 2], "actuators": [1]})), "demands")
    assert ("The net is held by 2 wires, an unrecorded number of column heads and 1 "
            "grabbed node.") in section
    missing = _sized_model()
    missing["held"] = None
    assert "Which nodes hold the net is not recorded." in _section(
        exports.datasheet_markdown(missing), "demands")
    # the first version keeps its anchors and its cut rule
    first = _section(exports.datasheet_markdown(_model()), "demands")
    assert "The net is held by 36 anchors and driven by 1 wire." in first
    assert "for the cut rule" in first


def test_a_version_2_document_assumes_a_prestress_floor_not_a_cut_rule():
    by_what = {a["what"]: a for a in _sized_model()["assumptions"]}
    assert by_what["uniform prestress"]["why"] == (
        "The prestress is the floor every member is given; the sag is judged at it.")
    assert by_what["rope EA"]["why"].endswith(
        "The net's stiffness, through which the sag is worked out, scales with it, so a "
        "different EA means a different sag.")
    text = exports.datasheet_markdown(_sized_model())
    assert "cut lengths" not in _section(text, "assumptions")
    assert "The cut rule" not in text
    first = {a["what"]: a for a in _model()["assumptions"]}
    assert first["uniform prestress"]["why"] == (
        "The cut rule assumes one uniform prestress across the net.")
    assert first["rope EA"]["why"].endswith("so a different EA means different cut lengths.")


def test_what_binds_is_said_at_the_sizing_stage_in_a_version_2_document():
    model = _sized_model()
    detail = model["capacity"]["detail"]
    assert detail and not detail.startswith("at the sizing stage")
    assert _chosen_row(model, "Binds on")[2] == "at the sizing stage: {}".format(detail)
    # the engine's sentence itself is unchanged, and so is the first version's cell
    import catalogue
    first = _model(row=dict(_row(catalogue.load_parts(), _configuration()),
                            load_factor={"binding_part": "x", "detail": "y"}))
    assert _chosen_row(first, "Binds on")[2] == "y"


# ---------------------------------------------------------------------------
# The acceptance line is a tolerance from the designed form (the owner's ruling
# of 9 October 2026). The data sheet says why the machine is not judged against
# the mould it replaces, then the tolerance, then the run's notes; the tolerance
# is listed with the assumptions; a document from before keeps its line in its
# own recorded words.
# ---------------------------------------------------------------------------

_TOLERANCE_SOURCE = "a tolerance of 20 mm from the designed form, set for this run"
# the real study's line before the ruling: a catalogue rib nobody designed
_GLULAM_SOURCE = (
    'falsework glulam-rib-9000: {"depth": 400.0, "description": "glulam GL24h rib 9000 x '
    '600, 400 x 90 deep", "e_modulus": 11600.0, "spacing": 600.0, "span": 9000.0, '
    '"width": 90.0}')
_FORMWORK_NOTE = ("no formwork document, so no frames and no column heads: the net is "
                  "analysed at the finished shape held by its drum ends alone")


def test_the_first_section_says_why_a_tolerance_then_the_tolerance_then_the_notes():
    model = _sized_model(acceptance=20.0, acceptance_source=_TOLERANCE_SOURCE,
                         tolerance_mm=20.0, note=_FORMWORK_NOTE)
    section = _section(exports.datasheet_markdown(model), "replaces")
    why = section.index("A mould like that barely moves, so the machine is not judged "
                        "against it.")
    tolerance = section.index("This run holds the net to a tolerance of 20 mm from the "
                              "designed form, set for this run, and that tolerance is the "
                              "acceptance line.")
    note = section.index("The analysis carries this note, quoted verbatim: \"{}\".".format(
        _FORMWORK_NOTE))
    assert why < tolerance < note
    import re
    assert not re.search(r"\brib\b|catalogue", section), "a tolerance names no rib"
    # a tolerance that is not a whole millimetre keeps its figure
    half = _sized_model(acceptance=12.5, tolerance_mm=12.5, acceptance_source=(
        "a tolerance of 12.5 mm from the designed form, set for this run"))
    assert "a tolerance of 12.5 mm from the designed form" in _section(
        exports.datasheet_markdown(half), "replaces")


def test_a_document_from_before_the_tolerance_keeps_its_line_in_its_own_words():
    # the real study's document of 8 October: its line came from a rib, in the
    # engine's own words
    older = _sized_model(acceptance=3.25, acceptance_source=_GLULAM_SOURCE)
    section = _section(exports.datasheet_markdown(older), "replaces")
    assert ("This document was written before the tolerance: its acceptance line of "
            "3.25 mm was recorded in these words, quoted verbatim: \"{}\". It is taken "
            "from that record and is not re-derived here.".format(_GLULAM_SOURCE)) in section
    # a line with no source, and no line at all, are each said plainly
    unsourced = _section(exports.datasheet_markdown(
        _sized_model(acceptance=3.25, acceptance_source=None)), "replaces")
    assert ("The acceptance line is 3.25 mm. No source was recorded for it, so the reader "
            "should treat the line with caution.") in unsourced
    lineless = _section(exports.datasheet_markdown(
        _sized_model(acceptance=None, acceptance_source=None)), "replaces")
    assert "No acceptance line is set for this run, so the net's shape is not judged." in lineless


def test_the_tolerance_is_listed_under_the_assumptions_in_both_documents():
    model = _sized_model(acceptance=20.0, acceptance_source=_TOLERANCE_SOURCE,
                         tolerance_mm=20.0)
    entry = next(a for a in model["assumptions"] if a["what"] == "acceptance tolerance")
    assert entry["value"] == 20.0 and entry["unit"] == "mm"
    assert entry["why"] == ("The most the net may stray from its designed form, set for this "
                            "run: the line its sag is judged against.")
    sheet = _section(exports.datasheet_markdown(model), "assumptions")
    assert ("**acceptance tolerance.** Value: 20 mm. The most the net may stray from its "
            "designed form, set for this run: the line its sag is judged against.") in sheet
    rows = exports._sheet_rows(model)["Read this"]
    assert any(row and row[0] == "acceptance tolerance" for row in rows)
    # the rib the server once swapped in is no assumption of any document now,
    # and a document with no tolerance lists none
    older = _sized_model(acceptance=3.25, acceptance_source=_GLULAM_SOURCE,
                         note=_FORMWORK_NOTE + "; falsework glulam-rib-9000 was used in "
                              "place of plywood-rib-2000, which spans 2000 mm, less than "
                              "half the vault's 16.1 m reach")
    assert not any("falsework" in a["what"] or a["what"] == "acceptance tolerance"
                   for a in older["assumptions"])
